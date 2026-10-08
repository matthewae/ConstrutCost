"""KF-13 / KF-17: export RAB dengan kolom pilihan, rekap biaya, RAB rinci, dan kebutuhan besi."""

import shutil
import subprocess

import pytest

from conftest import DUPLEX


@pytest.fixture
def data_export(db_sementara):
    from database.biaya_repository import tambah_biaya
    from database.proyek_repository import create_proyek
    from estimasi_service import jalankan_estimasi
    from export_service import ambil_data_export

    pid = create_proyek("Duplex", str(DUPLEX))
    jalankan_estimasi(pid)
    tambah_biaya(pid, "Biaya perencanaan teknis", "persen", 3)
    tambah_biaya(pid, "Biaya perizinan (PBG / SLF)", "nilai", 7_500_000)
    return ambil_data_export(pid)


META = {"nama_proyek": "Rumah Uji", "lokasi": "Kota Bandung", "pemilik": "Bpk. Uji", "tahun": 2027}


def _header(ws, teks="Uraian Pekerjaan"):
    for row in ws.iter_rows():
        nilai = [c.value for c in row]
        if teks in nilai:
            return [v for v in nilai if v is not None]
    raise AssertionError("header tidak ditemukan")


def test_kolom_dan_sheet_sesuai_pilihan(data_export, tmp_path):
    from openpyxl import load_workbook

    from export_service import OpsiExport, export_excel

    lengkap = tmp_path / "lengkap.xlsx"
    export_excel(lengkap, data_export, META, OpsiExport(kolom=("no", "kode", "harga", "bobot", "rumus")))
    wb = load_workbook(lengkap)
    lantai = [n for n in wb.sheetnames if n.startswith("Lt")]
    assert lantai and wb.sheetnames == ["RAB", "Rekapitulasi", "RAB Rinci", "Kebutuhan Besi", "Rekap per Lantai",
                                        "Kebutuhan per Lantai", *lantai, "Detail Elemen"]
    assert _header(wb["RAB"]) == [
        "No", "Kode Analisa", "Uraian Pekerjaan", "Volume", "Sat", "Harga Satuan (Rp)", "Jumlah Harga (Rp)", "Bobot",
    ]
    assert "Uraian Rumus" in _header(wb["Detail Elemen"])

    minimal = tmp_path / "minimal.xlsx"
    export_excel(minimal, data_export, META, OpsiExport(kolom=(), rekap=False, rinci=False, besi=False, detail=False, lantai=False,
                                                         per_lantai=False))
    wb = load_workbook(minimal)
    assert wb.sheetnames == ["RAB"]
    # kolom wajib selalu ada
    assert _header(wb["RAB"]) == ["Uraian Pekerjaan", "Volume", "Sat", "Jumlah Harga (Rp)"]


def test_rab_rinci_per_tipe_elemen(data_export, tmp_path):
    from openpyxl import load_workbook

    from export_service import OpsiExport, export_excel

    path = tmp_path / "rab.xlsx"
    export_excel(path, data_export, META, OpsiExport())
    teks = [c.value for row in load_workbook(path)["RAB Rinci"].iter_rows() for c in row if isinstance(c.value, str)]
    assert "Balok B1 (18/41) — 6 buah" in teks
    assert any("Tulangan utama 6 D16" in t for t in teks)
    assert any("Sengkang Ø10-150" in t for t in teks)
    besi = [c.value for row in load_workbook(path)["Kebutuhan Besi"].iter_rows() for c in row]
    assert {"Ø10", "D13", "D16"} <= set(besi)


def test_rekap_biaya_tidak_langsung_dan_terbilang(data_export, tmp_path):
    from openpyxl import load_workbook

    from export_service import OpsiExport, export_excel

    path = tmp_path / "rab.xlsx"
    export_excel(path, data_export, META, OpsiExport())
    teks = [c.value for row in load_workbook(path)["Rekapitulasi"].iter_rows() for c in row if isinstance(c.value, str)]
    assert "BIAYA TIDAK LANGSUNG" in teks
    assert "Biaya perencanaan teknis (3% × A)" in teks
    assert "TOTAL DIBULATKAN" in teks
    assert any(t.startswith("Terbilang: ") and t.endswith("rupiah") for t in teks)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice tidak tersedia")
def test_rumus_excel_dihitung_tanpa_galat(data_export, tmp_path):
    """Hitung ulang semua rumus dengan LibreOffice: tidak ada #REF!/#NAME?, total sama dengan aplikasi,
    dan selisih sheet rinci & detail terhadap sheet RAB = 0."""
    from openpyxl import load_workbook

    from export_service import OpsiExport, export_excel

    sumber = tmp_path / "rab.xlsx"
    export_excel(sumber, data_export, META, OpsiExport(kolom=("no", "kode", "harga", "bobot")))
    keluar = tmp_path / "hitung"
    subprocess.run(
        ["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(keluar), str(sumber)],
        check=True, capture_output=True, timeout=180,
    )
    wb = load_workbook(keluar / "rab.xlsx", data_only=True)
    ring = data_export["ringkasan"]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert not (isinstance(c.value, str) and c.value.startswith(("#", "Err:"))), (ws.title, c.coordinate)
            label = next((c.value for c in row if isinstance(c.value, str)), "")
            angka = [c.value for c in row if isinstance(c.value, (int, float))]
            if label == "TOTAL DIBULATKAN":
                assert angka[-1] == ring["dibulatkan"]
            if label.startswith("Selisih terhadap sheet RAB"):
                assert angka[-1] == pytest.approx(0, abs=0.01)


def test_pdf_portrait_dan_landscape(data_export, tmp_path):
    from pypdf import PdfReader

    from export_service import OpsiExport, export_pdf

    tegak = tmp_path / "tegak.pdf"
    datar = tmp_path / "datar.pdf"
    export_pdf(tegak, data_export, META, OpsiExport(kolom=("no", "harga")))
    export_pdf(datar, data_export, META, OpsiExport(kolom=(), orientasi_pdf="landscape", detail=False))
    a, b = PdfReader(tegak), PdfReader(datar)
    assert a.pages[0].mediabox.width < a.pages[0].mediabox.height
    assert b.pages[0].mediabox.width > b.pages[0].mediabox.height
    isi = " ".join(p.extract_text() for p in a.pages)
    for teks in ("REKAPITULASI", "RAB RINCI PER TIPE ELEMEN", "KEBUTUHAN BESI", "Terbilang", "Bpk. Uji"):
        assert teks in isi


def test_pilihan_export_diingat(db_sementara, tmp_path):
    from dataclasses import replace

    from database.preferensi_repository import muat_preferensi, simpan_pilihan_export

    simpan_pilihan_export(replace(muat_preferensi(), kolom_laporan="no,bobot", isi_besi=False,
                                  orientasi_pdf="landscape", direktori_export=str(tmp_path)))
    p = muat_preferensi()
    assert p.kolom == ("no", "bobot") and p.isi_besi is False and p.orientasi_pdf == "landscape"
