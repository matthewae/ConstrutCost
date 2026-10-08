"""Rincian kebutuhan per lantai (beton, bekisting, besi per diameter, bahan/tenaga/alat AHSP) dan
export satu sheet / bagian PDF per lantai untuk model dengan banyak lantai (lebih dari 2)."""

import os
import shutil
import subprocess

import pytest

from conftest import DUPLEX

N_LANTAI = 5


@pytest.fixture(scope="module")
def ifc_bertingkat(tmp_path_factory):
    """Model sintetis 5 lantai (fondasi + 3 lantai + dak). Setiap lantai: 2 kolom 25/25, 1 balok 20/40,
    1 pelat 12 cm, dan 1 dinding. Nama lantai sengaja memuat karakter yang tidak boleh di nama sheet."""
    import ifcopenshell.api as api

    f = api.run("project.create_file", version="IFC4")
    proj = api.run("root.create_entity", f, ifc_class="IfcProject", name="Rumah Bertingkat")
    api.run("unit.assign_unit", f, length={"is_metric": True, "raw": "METERS"})  # profil kolom dalam meter
    model = api.run("context.add_context", f, context_type="Model")
    body = api.run("context.add_context", f, context_type="Model", context_identifier="Body",
                   target_view="MODEL_VIEW", parent=model)
    site = api.run("root.create_entity", f, ifc_class="IfcSite")
    gedung = api.run("root.create_entity", f, ifc_class="IfcBuilding")
    api.run("aggregate.assign_object", f, products=[site], relating_object=proj)
    api.run("aggregate.assign_object", f, products=[gedung], relating_object=site)
    nama = ["00 FONDASI", "01 LANTAI 1", "02 LANTAI 2", "03 LANTAI 3/MEZZANINE", "04 DAK [ATAP]"]
    for i, nm in enumerate(nama):
        lt = api.run("root.create_entity", f, ifc_class="IfcBuildingStorey", name=nm)
        lt.Elevation = -1.5 + 3.5 * i if i else -1.5
        api.run("aggregate.assign_object", f, products=[lt], relating_object=gedung)
        produk = []
        for k in range(2):
            kol = api.run("root.create_entity", f, ifc_class="IfcColumn", name=f"K{i}{k}")
            api.run("geometry.assign_representation", f, product=kol, representation=api.run(
                "geometry.add_profile_representation", f, context=body,
                profile=f.createIfcRectangleProfileDef("AREA", None, None, 0.25, 0.25), depth=3.5))
            produk.append(kol)
        balok = api.run("root.create_entity", f, ifc_class="IfcBeam", name=f"B{i}")
        api.run("geometry.assign_representation", f, product=balok, representation=api.run(
            "geometry.add_profile_representation", f, context=body,
            profile=f.createIfcRectangleProfileDef("AREA", None, None, 0.2, 0.4), depth=4.0))
        pelat = api.run("root.create_entity", f, ifc_class="IfcSlab", name=f"P{i}", predefined_type="FLOOR")
        api.run("geometry.assign_representation", f, product=pelat, representation=api.run(
            "geometry.add_slab_representation", f, context=body, depth=0.12,
            polyline=[(0.0, 0.0), (4.0, 0.0), (4.0, 4.0), (0.0, 4.0)]))
        dinding = api.run("root.create_entity", f, ifc_class="IfcWall", name=f"D{i}")
        api.run("geometry.assign_representation", f, product=dinding, representation=api.run(
            "geometry.add_wall_representation", f, context=body, length=4, height=3, thickness=0.15))
        api.run("spatial.assign_container", f, products=produk + [balok, pelat, dinding], relating_structure=lt)
    path = tmp_path_factory.mktemp("ifc") / "bertingkat.ifc"
    f.write(str(path))
    return path


@pytest.fixture
def proyek_bertingkat(db_sementara, ifc_bertingkat):
    from database.proyek_repository import create_proyek
    from database.seed_data import seed_pekerjaan
    from estimasi_service import jalankan_estimasi
    from export_service import ambil_data_export

    seed_pekerjaan()
    pid = create_proyek("Rumah 5 Lantai", str(ifc_bertingkat))
    jalankan_estimasi(pid)
    return pid, ambil_data_export(pid)


# ---------------------------------------------------------------- data


def test_lantai_dikenali_semua_urut_elevasi(proyek_bertingkat):
    from kebutuhan_lantai import rincian_per_lantai

    _, d = proyek_bertingkat
    rl = rincian_per_lantai(d["baris"], d["komponen"])
    assert [x["lantai"] for x in rl] == [
        "00 FONDASI", "01 LANTAI 1", "02 LANTAI 2", "03 LANTAI 3/MEZZANINE", "04 DAK [ATAP]"]
    assert sum(x["total"] for x in rl) == pytest.approx(d["ringkasan"]["langsung"], abs=0.01)
    for x in rl:
        assert x["jumlah_elemen"] == 5
        assert sum(k["total"] for k in x["rinci"]) == pytest.approx(x["total"], abs=0.01)


def test_struktur_besi_per_diameter_konsisten(proyek_bertingkat):
    from kebutuhan_lantai import matriks_lantai, rincian_per_lantai, ringkas_lantai

    _, d = proyek_bertingkat
    rl = rincian_per_lantai(d["baris"], d["komponen"])
    for x in rl:
        s = ringkas_lantai(x)
        # beton 2 kolom 0,25 x 0,25 x 3,5 + balok 0,2 x 0,4 x 4 + pelat 4 x 4 x 0,12 (dibulatkan geometri)
        assert s["beton"] > 2.0 and s["bekisting"] > 0 and s["besi"] > 0
        kolom = next(g for g in x["struktur"] if g["label"].startswith("Kolom"))
        assert kolom["jumlah_elemen"] == 2 and kolom["mutu"] == "f'c 20 MPa"
        assert 40 < kolom["rasio"] < 350
        # besi per diameter + asumsi rasio = semua pembesian lantai itu
        assert sum(k["berat"] for k in x["besi"]) + x["besi_rasio"] == pytest.approx(s["besi"])
        for k in x["besi"]:
            assert k["batang"] >= k["panjang"] / 12
    # tabel silang per diameter menjumlah ke kebutuhan besi proyek
    total_d = {f"{k['label']} {k['jenis']}": k["berat"] for k in d["besi"]}
    for m in matriks_lantai(rl, "besi"):
        if m["label"] in total_d:
            assert m["total"] == pytest.approx(total_d[m["label"]])
        assert len(m["per_lantai"]) == N_LANTAI


def test_bahan_tenaga_alat_sesuai_koefisien_ahsp(proyek_bertingkat):
    from database.estimasi_repository import BUK_RATE
    from kebutuhan_lantai import rincian_per_lantai

    _, d = proyek_bertingkat
    rl = rincian_per_lantai(d["baris"], d["komponen"])
    biaya = sum(i["biaya"] for x in rl for t in ("bahan", "upah", "alat") for i in x["sumber_daya"][t])
    # tanpa harga khusus, harga dasar sumber daya x (1 + BUK) = biaya langsung
    assert biaya * (1 + BUK_RATE) == pytest.approx(d["ringkasan"]["langsung"], rel=1e-9)
    bahan = {i["nama"]: i for i in rl[1]["sumber_daya"]["bahan"]}
    semen = bahan["Semen Portland (PC)"]
    assert "zak @ 50 kg" in semen["keterangan"]
    assert any("Pekerja" == i["nama"] for i in rl[1]["sumber_daya"]["upah"])


def test_varian_sumber_daya_digabung():
    from kebutuhan_lantai import keterangan_sumber_daya, nama_dasar

    sheet = {"Galian Tanah", "Beton"}
    assert nama_dasar("Pekerja (Galian Tanah)", sheet) == "Pekerja"
    assert nama_dasar("Pasir beton (Beton)", sheet) == "Pasir beton"
    assert nama_dasar("Pasir pasang (quarry - lokasi pekerjaan)", sheet) == "Pasir pasang (quarry - lokasi pekerjaan)"
    assert keterangan_sumber_daya("Semen Portland (PC)", "kg", 1001) == "≈ 21 zak @ 50 kg"
    assert keterangan_sumber_daya("Pekerja", "OH", 3) == "orang-hari"


def test_nama_sheet_lantai_aman_dan_unik():
    from export_service import nama_sheet_lantai

    pakai = {"rab", "rekap per lantai"}
    a = nama_sheet_lantai(1, "03 LANTAI 3/MEZZANINE [baru]: *?", pakai)
    assert not set("[]:*?/\\") & set(a) and len(a) <= 31 and a.startswith("Lt01 ")
    panjang = "Lantai dengan nama yang sangat panjang sekali dari Revit"
    b, c = nama_sheet_lantai(2, panjang, pakai), nama_sheet_lantai(2, panjang, pakai)
    assert b != c and len(b) <= 31 and len(c) <= 31
    assert nama_sheet_lantai(12, "Roof", pakai) == "Lt12 Roof"


# ---------------------------------------------------------------- export


def test_export_excel_satu_sheet_per_lantai(proyek_bertingkat, tmp_path):
    from openpyxl import load_workbook

    from export_service import OpsiExport, export_excel

    _, d = proyek_bertingkat
    path = export_excel(tmp_path / "rab.xlsx", d, {}, OpsiExport())
    wb = load_workbook(path)
    lembar = [n for n in wb.sheetnames if n.startswith("Lt")]
    assert len(lembar) == N_LANTAI
    assert lembar[3] == "Lt04 03 LANTAI 3 MEZZANINE" and lembar[4] == "Lt05 04 DAK ATAP"
    assert "Kebutuhan per Lantai" in wb.sheetnames and "Rekap per Lantai" in wb.sheetnames
    ws = wb[lembar[1]]
    teks = [c.value for c in ws["B"] if isinstance(c.value, str)]
    for judul in ("STRUKTUR BETON PER TIPE ELEMEN", "KEBUTUHAN BESI TULANGAN PER DIAMETER",
                  "KEBUTUHAN BAHAN / MATERIAL (KOEFISIEN AHSP)", "KEBUTUHAN TENAGA KERJA (KOEFISIEN AHSP)",
                  "RINCIAN PEKERJAAN (RAB RINCI LANTAI)"):
        assert judul in teks
    wk = wb["Kebutuhan per Lantai"]
    judul_lantai = [c.value for c in next(r for r in wk.iter_rows() if r[0].value == "No")]
    assert judul_lantai[3:3 + N_LANTAI] == [
        "00 FONDASI", "01 LANTAI 1", "02 LANTAI 2", "03 LANTAI 3/MEZZANINE", "04 DAK [ATAP]"]
    # opsi dimatikan: tidak ada sheet per lantai
    path2 = export_excel(tmp_path / "rab2.xlsx", d, {}, OpsiExport(per_lantai=False, lantai=False))
    assert not [n for n in load_workbook(path2).sheetnames if n.startswith("Lt") or "per Lantai" in n]


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice tidak tersedia")
def test_export_excel_per_lantai_rumus_cocok(proyek_bertingkat, tmp_path):
    """Dihitung ulang LibreOffice: tidak ada #ERROR dan semua baris 'Selisih' bernilai 0."""
    from openpyxl import load_workbook

    from export_service import OpsiExport, export_excel

    _, d = proyek_bertingkat
    sumber = export_excel(tmp_path / "rab.xlsx", d, {}, OpsiExport(kolom=("no", "kode", "harga", "bobot")))
    keluar = tmp_path / "hitung"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(keluar), sumber],
                   check=True, capture_output=True, timeout=180, env=dict(os.environ, HOME=str(tmp_path)))
    wb = load_workbook(keluar / "rab.xlsx", data_only=True)
    selisih = 0
    for ws in wb.worksheets:
        for baris in ws.iter_rows():
            for c in baris:
                assert not (isinstance(c.value, str) and c.value.startswith("#")), f"{ws.title}!{c.coordinate} {c.value}"
            if any(isinstance(c.value, str) and c.value.startswith("Selisih") for c in baris):
                angka = [c.value for c in baris if isinstance(c.value, (int, float))]
                assert angka and abs(angka[-1]) < 0.01, ws.title
                selisih += 1
    assert selisih >= N_LANTAI + 2
    wl = wb["Rekap per Lantai"]
    total = next(r for r in wl.iter_rows() if r[1].value == f"JUMLAH {N_LANTAI} LANTAI")
    assert total[6].value == pytest.approx(d["ringkasan"]["langsung"], abs=0.5)


def test_export_pdf_bagian_per_lantai(proyek_bertingkat, tmp_path):
    from pypdf import PdfReader

    from export_service import OpsiExport, export_pdf

    _, d = proyek_bertingkat
    path = export_pdf(tmp_path / "rab.pdf", d, {}, OpsiExport(orientasi_pdf="landscape"))
    teks = " ".join(p.extract_text() for p in PdfReader(path).pages)
    for i in range(1, N_LANTAI + 1):
        assert f"RINCIAN KEBUTUHAN LANTAI {i} DARI {N_LANTAI}" in teks
    assert "Kebutuhan besi per diameter per lantai" in teks and "REKAPITULASI BIAYA & KEBUTUHAN STRUKTUR PER LANTAI" in teks


def test_export_duplex_empat_lantai(db_sementara, tmp_path):
    """Model nyata Revit (IFC2x3) dengan 4 lantai termasuk fondasi dan atap."""
    from openpyxl import load_workbook

    from database.proyek_repository import create_proyek
    from database.seed_data import seed_pekerjaan
    from estimasi_service import jalankan_estimasi
    from export_service import OpsiExport, ambil_data_export, export_excel

    seed_pekerjaan()
    pid = create_proyek("Duplex", str(DUPLEX))
    jalankan_estimasi(pid)
    path = export_excel(tmp_path / "d.xlsx", ambil_data_export(pid), {}, OpsiExport())
    assert [n for n in load_workbook(path).sheetnames if n.startswith("Lt")] == [
        "Lt01 T FDN", "Lt02 Level 1", "Lt03 Level 2", "Lt04 Roof"]


# ---------------------------------------------------------------- tampilan


def test_halaman_per_lantai_ringkas_dan_rinci(proyek_bertingkat):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication

    test_halaman_per_lantai_ringkas_dan_rinci.app = QApplication.instance() or QApplication([])
    from gui.estimasi_page import SEMUA_LANTAI, EstimasiPage

    pid, _ = proyek_bertingkat
    hal = EstimasiPage(pid, "Rumah 5 Lantai")
    hal.muat()
    hal.grup_mode.button(hal.MODE_LANTAI).setChecked(True)
    hal._isi_ulang()
    t = hal.tabel
    judul = [t.horizontalHeaderItem(i).text() for i in range(t.columnCount())]
    assert judul[3:6] == ["BETON (m³)", "BEKISTING (m²)", "BESI (kg)"]
    lantai = [t.item(r, 1).text() for r in range(t.rowCount())
              if t.item(r, 0) and t.item(r, 0).text().isdigit()]
    assert len(lantai) == N_LANTAI and lantai[0].startswith("00 FONDASI")
    assert t.item(t.rowCount() - 1, 1).text() == f"JUMLAH {N_LANTAI} LANTAI"

    hal._klik_ganda(0, 1)  # klik dua kali lantai pertama -> rincian lantai itu
    assert hal.combo_lantai.currentText() == "00 FONDASI"
    teks = [t.item(r, 1).text() for r in range(t.rowCount()) if t.item(r, 1)]
    for judul in ("STRUKTUR BETON PER TIPE ELEMEN", "KEBUTUHAN BESI TULANGAN PER DIAMETER",
                  "KEBUTUHAN BAHAN / MATERIAL (KOEFISIEN AHSP)", "RINCIAN PEKERJAAN (RAB RINCI LANTAI)"):
        assert judul in teks
    assert any(x.startswith("Kolom") for x in teks)
    hal.combo_lantai.setCurrentText(SEMUA_LANTAI)
    assert t.horizontalHeaderItem(3).text() == "BETON (m³)"


# ---------------------------------------------------------------- rincian perhitungan & penutup bangunan


def _teks_panel(panel) -> list:
    from PySide6.QtWidgets import QLabel

    return [lb.text() for lb in panel._isi.findChildren(QLabel)]


def test_penutup_bangunan_atap_dak_plafon(db_sementara, proyek_bertingkat):
    from conftest import AC20
    from database.proyek_repository import create_proyek
    from estimasi_service import jalankan_estimasi
    from export_service import ambil_data_export
    from kebutuhan_lantai import penutup_bangunan, rincian_per_lantai

    _, d = proyek_bertingkat
    # model sintetis tanpa atap/dak: penutupnya plafon di bawah pelat lantai teratas
    lantai, unsur = penutup_bangunan(rincian_per_lantai(d["baris"], d["komponen"]))
    assert lantai == "04 DAK [ATAP]" and [j for j, _ in unsur] == ["Plafon"]
    assert penutup_bangunan([]) == (None, [])
    hasil = {}
    for nama, path in (("AC20", AC20), ("Duplex", DUPLEX)):
        pid = create_proyek(nama, str(path))
        jalankan_estimasi(pid)
        dd = ambil_data_export(pid)
        hasil[nama] = penutup_bangunan(rincian_per_lantai(dd["baris"], dd["komponen"]))
    lantai, unsur = hasil["AC20"]  # atap genteng + rangka baja ringan, plafon gypsum di lantai teratas
    assert lantai == "Dachgeschoss" and [j for j, _ in unsur] == ["Atap", "Plafon"]
    assert "Genteng" in dict(unsur)["Atap"] and "Gypsum" in dict(unsur)["Plafon"]
    lantai, unsur = hasil["Duplex"]  # dak beton di lantai Roof
    assert lantai == "Roof" and unsur[0][0] == "Dak beton"


def test_panel_rincian_perhitungan_per_lantai(proyek_bertingkat):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication

    test_panel_rincian_perhitungan_per_lantai.app = QApplication.instance() or QApplication([])
    from gui.estimasi_page import EstimasiPage

    pid, _ = proyek_bertingkat
    hal = EstimasiPage(pid, "Rumah 5 Lantai")
    hal.muat()
    hal.grup_mode.button(hal.MODE_LANTAI).setChecked(True)
    hal._isi_ulang()
    teks = _teks_panel(hal.panel)  # tanpa baris dipilih: semua lantai berurutan
    assert teks[0] == f"RINCIAN PERHITUNGAN PER LANTAI ({N_LANTAI} LANTAI)"
    judul = [t for t in teks if t in ("00 FONDASI", "01 LANTAI 1", "02 LANTAI 2", "03 LANTAI 3/MEZZANINE", "04 DAK [ATAP]")]
    assert judul == ["00 FONDASI", "01 LANTAI 1", "02 LANTAI 2", "03 LANTAI 3/MEZZANINE", "04 DAK [ATAP]"]
    assert any(t.startswith("Besi D") or t.startswith("Besi Ø") for t in teks)
    assert any("batang @ 12 m" in t for t in teks) and any("zak @ 50 kg" in t for t in teks)
    assert "PENUTUP BANGUNAN" in teks and "Plafon" in teks
    assert teks.index("PENUTUP BANGUNAN") > teks.index("04 DAK [ATAP]")  # penutup ditulis paling akhir

    hal.tabel.selectRow(0)  # pilih satu lantai: rincian lengkap lantai itu saja
    teks = _teks_panel(hal.panel)
    assert teks[0] == "RINCIAN PERHITUNGAN LANTAI" and "00 FONDASI" in teks and "01 LANTAI 1" not in teks
    assert "TENAGA KERJA (VOLUME × KOEFISIEN AHSP)" in teks
    assert any(t.startswith("• ") and " bh × " in t for t in teks)  # backup volume tiap item
    hal.tabel.clearSelection()
    assert _teks_panel(hal.panel)[0].startswith("RINCIAN PERHITUNGAN PER LANTAI")

    hal.grup_mode.button(hal.MODE_REKAP).setChecked(True)
    hal._isi_ulang()
    assert _teks_panel(hal.panel)[0] == "RINCIAN PERHITUNGAN"  # mode lain: panel bawaan


def test_sheet_rincian_perhitungan_urut_sampai_penutup(db_sementara, tmp_path):
    from openpyxl import load_workbook

    from conftest import AC20
    from database.proyek_repository import create_proyek
    from estimasi_service import jalankan_estimasi
    from export_service import OpsiExport, ambil_data_export, export_excel

    pid = create_proyek("AC20", str(AC20))
    jalankan_estimasi(pid)
    path = export_excel(tmp_path / "r.xlsx", ambil_data_export(pid), {}, OpsiExport())
    ws = load_workbook(path)["Rincian Perhitungan Lantai"]
    b = [c.value for c in ws["B"] if isinstance(c.value, str)]
    assert b.index("ERDGESCHOSS") < b.index("DACHGESCHOSS") < b.index("PENUTUP BANGUNAN (Dachgeschoss)")
    assert b[-2:] == ["Atap", "Plafon"]
    assert "PEKERJAAN TANAH" in b and "PEKERJAAN BETON" in b and "PEKERJAAN KUDA-KUDA DAN ATAP" in b
    assert any(x.startswith("KEBUTUHAN BESI TULANGAN PER DIAMETER") for x in b)
    assert any(x.startswith("KEBUTUHAN BAHAN / MATERIAL") for x in b)
    assert any(x.strip().startswith("• ") for x in b)  # baris backup volume
    # kolom backup sesuai lampiran RAB: Panjang | Lebar | Tinggi | Luas | Jlh Unit | Volume | Total Volume
    judul = [c.value for c in next(r for r in ws.iter_rows() if r[0].value == "No")]
    assert judul[2:9] == ["Panjang (m)", "Lebar (m)", "Tinggi / Tebal (m)", "Luas (m²)", "Jlh Unit", "Volume / Unit",
                          "Total Volume"]
    item = next(r for r in ws.iter_rows() if isinstance(r[8].value, str) and r[8].value.startswith("=SUM(I"))
    assert item[11].value == f"=I{item[0].row}*K{item[0].row}"  # jumlah = total volume x harga satuan
    m = [c.value for c in ws["M"] if isinstance(c.value, str)]
    assert any("batang @ 12 m" in x for x in m) and any("zak @ 50 kg" in x for x in m)


def test_rab_lantai_susunan_rap_dan_backup(proyek_bertingkat):
    from kebutuhan_lantai import judul_kategori, rab_lantai, rincian_per_lantai

    _, d = proyek_bertingkat
    for x in rincian_per_lantai(d["baris"], d["komponen"]):
        bagian = rab_lantai(x)
        assert [b["no"] for b in bagian] == ["I.", "II.", "III.", "IV.", "V.", "VI.", "VII.", "VIII.", "IX.", "X."][:len(bagian)]
        assert all(b["judul"] == judul_kategori(b["kategori"]) for b in bagian)
        assert sum(b["total"] for b in bagian) == pytest.approx(x["total"])
        beton = next(b for b in bagian if b["kategori"] == "Beton")
        kolom = next(g for g in beton["grup"] if g["judul"] and g["judul"].startswith("Kolom"))
        assert kolom["judul"].endswith("— 2 buah") and kolom["no"] == "1"
        uraian = [it["uraian"] for it in kolom["items"]]
        assert uraian[0].startswith("Beton Kolom") and any(u.startswith("Tulangan utama") for u in uraian)
        for b in bagian:
            for g in b["grup"]:
                for it in g["items"]:
                    assert it["no"] == "-" if g["judul"] else it["no"].isdigit()
                    assert sum(e["volume"] for e in it["backup"]) == pytest.approx(it["volume"])
                    for e in it["backup"]:
                        assert e["unit"] * e["per_unit"] == pytest.approx(e["volume"])
        backup = kolom["items"][0]["backup"]  # 2 kolom 25/25 tinggi 3,5 m (nama berbeda: K..0, K..1)
        assert sum(e["unit"] for e in backup) == 2
        assert all(e["per_unit"] == pytest.approx(0.25 * 0.25 * 3.5, rel=1e-3) for e in backup)
        assert all(e["tinggi"] == pytest.approx(3.5) for e in backup)

    from kebutuhan_lantai import backup_volume

    sama = [{"nama_elemen": "Kolom:K1 20/25:3471", "panjang": None, "lebar": 0.2, "tinggi": 3.5, "luas": None,
             "volume_pekerjaan": 0.175, "rumus": "V = b × h × t"} for _ in range(3)]
    b = backup_volume(sama + [{**sama[0], "nama_elemen": "Kolom:K1 20/25:9999", "volume_pekerjaan": 0.17500001}])
    assert len(b) == 1 and b[0]["uraian"] == "Kolom:K1 20/25" and b[0]["unit"] == 4  # id Revit dibuang, digabung


def test_tabel_per_lantai_judul_bagian_dan_panel_kategori(proyek_bertingkat):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QApplication

    test_tabel_per_lantai_judul_bagian_dan_panel_kategori.app = QApplication.instance() or QApplication([])
    from gui.estimasi_page import EstimasiPage

    pid, _ = proyek_bertingkat
    hal = EstimasiPage(pid, "Rumah 5 Lantai")
    hal.resize(900, 700)  # layar sempit: kolom judul tidak boleh menyempit
    hal.show()
    hal.muat()
    hal.grup_mode.button(hal.MODE_LANTAI).setChecked(True)
    hal._isi_ulang()
    QApplication.processEvents()
    t = hal.tabel
    assert t.columnWidth(1) >= 300
    judul = [t.item(r, 1).text().strip() for r in range(t.rowCount()) if t.item(r, 1)]
    assert "I.  PEKERJAAN BETON" in judul or any(j.endswith("PEKERJAAN BETON") for j in judul)
    baris_beton = next(r for r in range(t.rowCount()) if t.item(r, 0) and (t.item(r, 0).data(Qt.UserRole) or ("",))[0]
                       == "kategori" and t.item(r, 0).data(Qt.UserRole)[1] == ("01 LANTAI 1", "Beton"))
    t.selectRow(baris_beton)
    teks = _teks_panel(hal.panel)
    assert teks[0] == "RINCIAN PERHITUNGAN · PEKERJAAN BETON" and "01 LANTAI 1" in teks
    assert not any("PEKERJAAN DINDING" in x for x in teks)  # hanya bagian yang dipilih
    assert any(x.startswith("1   Kolom") for x in teks) and any(x.strip().startswith("–  Beton Kolom") for x in teks)
    assert any(x.startswith("Contoh rumus:") for x in teks)
    hal.close()


def test_dialog_export_tombol_selalu_terlihat(proyek_bertingkat):
    """Layar pendek (offscreen 800 px, mirip laptop dengan skala 125-150%): isi digulir, tombol tetap terlihat."""
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication

    app = QApplication.instance() or QApplication([])
    test_dialog_export_tombol_selalu_terlihat.app = app
    from gui.export_dialog import ExportDialog

    pid, _ = proyek_bertingkat
    d = ExportDialog(pid, "Rumah_Tipe90_2Lantai_Carport")
    d.show()
    app.processEvents()
    layar = d.screen().availableGeometry()
    assert d.height() <= layar.height()
    for tombol in (d.btn_export, d.btn_pratinjau):
        assert tombol.isVisible()
        bawah = tombol.mapTo(d, tombol.rect().bottomRight())
        assert 0 < bawah.y() <= d.height() and 0 < bawah.x() <= d.width()
    assert d.btn_export.text().startswith("Export")
    d.close()
