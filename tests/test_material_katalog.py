"""Pilihan material per pekerjaan (combo box + bawaan) dan katalog HSPK lengkap."""

import os

import pytest

from database.hspk_bandung_2027 import AHSP, BAB, SATUAN
from database.material import SLOT, kode_dasar, kode_varian, pilihan_lengkap, varian_terpilih
from database.seed_data import PEMETAAN


def _sama_satuan(a: str, b: str) -> bool:
    return a.replace("’", "'").lower() == b.replace("’", "'").lower()


@pytest.mark.parametrize("slot", list(SLOT))
def test_semua_pilihan_ada_di_hspk_dengan_satuan_sama(slot):
    s = SLOT[slot]
    satuan = PEMETAAN[slot][2]
    assert len({o.kunci for o in s.opsi}) == len(s.opsi)
    assert s.bawaan in {o.kunci for o in s.opsi}
    for o in s.opsi:
        assert o.kode_hspk in AHSP, o
        assert _sama_satuan(SATUAN[o.kode_hspk], satuan), (o, SATUAN[o.kode_hspk], satuan)
        assert PEMETAAN[kode_varian(slot, o.kunci)][3] == o.kode_hspk


def test_bawaan_dinding_bata_ringan_60x20x10():
    assert PEMETAAN["DND.BATA"][3] == "3.6.4.2"
    assert "60x20x10" in PEMETAAN["DND.BATA"][0]
    assert PEMETAAN["DND.BATA~merah"][3] == "3.6.1.8"
    assert kode_dasar("DND.BATA~merah") == "DND.BATA" and kode_dasar("BTN.KOLOM") == "BTN.KOLOM"
    assert pilihan_lengkap({"DND.BATA": "tidak-ada"})["DND.BATA"] == "ringan10"
    assert varian_terpilih({"DND.BATA": "merah", "BTN.KOLOM": "m20"}) == {"DND.BATA": "DND.BATA~merah"}


def test_katalog_hspk_lengkap_bab_1_sampai_10():
    assert len(AHSP) > 2700
    assert set(BAB) == {str(i) for i in range(1, 11)}
    assert {k.split(".")[0] for k in AHSP} == set(BAB)
    # contoh item di luar lingkup rule engine yang kini tersedia di katalog
    assert SATUAN["3.18.3.1"] == "buah"  # kloset duduk
    assert "3.8.23.1" in AHSP  # cat berbasis air (sheet PENGECATAN (BARU))
    assert any(k.startswith("5.1.") for k in AHSP)  # instalasi listrik
    assert not any(k.split(".")[0].split("@")[0].isdigit() and int(k.split(".")[0]) > 10 for k in AHSP)


def _proyek_ac20():
    from conftest import AC20
    from database.proyek_repository import create_proyek
    from estimasi_service import jalankan_estimasi

    pid = create_proyek("AC20", str(AC20))
    jalankan_estimasi(pid)
    return pid


def _baris(pid, awalan):
    from database.estimasi_repository import get_hasil_estimasi_by_proyek

    return [r for r in get_hasil_estimasi_by_proyek(pid) if kode_dasar(r["kode_ahsp"]) == awalan]


def test_ganti_pilihan_mengganti_pekerjaan_tanpa_mengubah_volume(db_sementara):
    from database.estimasi_repository import _connect, get_harga_satuan_pekerjaan
    from database.material_repository import muat_pilihan, simpan_pilihan
    from estimasi_service import hitung_ulang_dari_elemen

    pid = _proyek_ac20()
    dinding = _baris(pid, "DND.BATA")
    assert dinding and all(r["kode_ahsp"] == "DND.BATA" for r in dinding)  # bawaan: bata ringan
    luas = sum(r["volume_pekerjaan"] for r in dinding)
    # volume satu baris diedit manual (KF-6): harus tetap setelah material diganti
    conn = _connect()
    conn.execute("UPDATE hasil_estimasi SET volume_pekerjaan = 7.5, diedit_manual = 1 WHERE id = ?",
                 (dinding[0]["hasil_id"],))
    conn.commit()
    conn.close()

    hasil = simpan_pilihan(pid, {"DND.BATA": "merah", "BTN.PELAT": "rm25"})
    assert set(hasil["berubah"]) == {"DND.BATA", "BTN.PELAT"}
    assert muat_pilihan(pid)["DND.BATA"] == "merah"
    baru = _baris(pid, "DND.BATA")
    assert {r["kode_ahsp"] for r in baru} == {"DND.BATA~merah"}
    assert baru[0]["nama_pekerjaan"].startswith("Pasangan Dinding Bata Merah")
    harga = get_harga_satuan_pekerjaan(baru[0]["pekerjaan_id"])
    assert harga == pytest.approx(426435, abs=1)  # HSPK 3.6.1.8
    ubah = next(r for r in baru if r["hasil_id"] == dinding[0]["hasil_id"])
    assert ubah["volume_pekerjaan"] == 7.5 and ubah["subtotal_biaya"] == pytest.approx(7.5 * harga)
    kolom = _baris(pid, "BTN.PELAT")
    assert kolom and {r["kode_ahsp"] for r in kolom} == {"BTN.PELAT~rm25"}
    assert "Ready Mix" in kolom[0]["nama_pekerjaan"]

    # hitung ulang dari elemen tetap memakai pilihan proyek
    hitung_ulang_dari_elemen(pid)
    assert {r["kode_ahsp"] for r in _baris(pid, "DND.BATA")} == {"DND.BATA~merah"}
    assert sum(r["volume_pekerjaan"] for r in _baris(pid, "DND.BATA")) == pytest.approx(luas)

    # kembali ke bawaan: kode polos dan pilihan tidak disimpan
    simpan_pilihan(pid, {})
    assert {r["kode_ahsp"] for r in _baris(pid, "DND.BATA")} == {"DND.BATA"}
    conn = _connect()
    assert conn.execute("SELECT material FROM proyek WHERE id = ?", (pid,)).fetchone()[0] is None
    conn.close()


def test_pilihan_ikut_file_proyek_dan_duplikat(db_sementara, tmp_path):
    from database.berkas_proyek import buka_berkas, duplikat_proyek, simpan_berkas
    from database.material_repository import muat_pilihan, simpan_pilihan

    pid = _proyek_ac20()
    simpan_pilihan(pid, {"KRM.LANTAI": "gr60", "ATP.PENUTUP": "metal"})
    salinan = duplikat_proyek(pid, "AC20 salinan")
    assert muat_pilihan(salinan)["KRM.LANTAI"] == "gr60"
    assert {r["kode_ahsp"] for r in _baris(salinan, "KRM.LANTAI")} == {"KRM.LANTAI~gr60"}
    path = simpan_berkas(pid, tmp_path / "ac20.coststruct", sertakan_ifc=False)
    dibuka = buka_berkas(path)
    assert muat_pilihan(dibuka)["ATP.PENUTUP"] == "metal"


def test_pilihan_tidak_dikenal_ditolak(db_sementara):
    from database.material_repository import simpan_pilihan
    from database.proyek_repository import create_proyek

    pid = create_proyek("Kosong", None)
    with pytest.raises(ValueError):
        simpan_pilihan(pid, {"BUKAN.SLOT": "x"})


def test_material_khusus_satu_elemen(db_sementara):
    from database.material_repository import material_elemen, simpan_material_elemen, simpan_pilihan
    from estimasi_service import hitung_ulang_dari_elemen, jalankan_estimasi

    pid = _proyek_ac20()
    dinding = _baris(pid, "DND.BATA")
    eid = dinding[0]["elemen_id"]
    info = material_elemen(eid)
    assert info["DND.BATA"]["kunci"] == "ringan10" and not info["DND.BATA"]["khusus"]
    assert {"PLS.DINDING", "CAT.DINDING"} <= set(info)

    simpan_material_elemen(eid, {"DND.BATA": "merah"})
    kode = {r["elemen_id"]: r["kode_ahsp"] for r in _baris(pid, "DND.BATA")}
    assert kode[eid] == "DND.BATA~merah"
    assert {k for e, k in kode.items() if e != eid} == {"DND.BATA"}  # dinding lain tetap bata ringan
    assert material_elemen(eid)["DND.BATA"]["khusus"]

    # pilihan proyek berubah: elemen dengan pilihan sendiri tidak ikut
    simpan_pilihan(pid, {"DND.BATA": "batako"})
    kode = {r["elemen_id"]: r["kode_ahsp"] for r in _baris(pid, "DND.BATA")}
    assert kode[eid] == "DND.BATA~merah" and {k for e, k in kode.items() if e != eid} == {"DND.BATA~batako"}

    # hitung ulang dari elemen dan baca ulang IFC tetap memakai pilihan elemen
    hitung_ulang_dari_elemen(pid)
    assert {r["elemen_id"]: r["kode_ahsp"] for r in _baris(pid, "DND.BATA")}[eid] == "DND.BATA~merah"
    jalankan_estimasi(pid)
    assert "DND.BATA~merah" in {r["kode_ahsp"] for r in _baris(pid, "DND.BATA")}

    # pilihan sama dengan proyek = kembali mengikuti proyek
    eid2 = next(r["elemen_id"] for r in _baris(pid, "DND.BATA") if r["kode_ahsp"] == "DND.BATA~merah")
    simpan_material_elemen(eid2, {"DND.BATA": "batako"})
    assert not material_elemen(eid2)["DND.BATA"]["khusus"]


@pytest.fixture(scope="module")
def qapp():
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication

    return QApplication.instance() or QApplication([])


def _teks_panel(panel) -> str:
    from PySide6.QtWidgets import QLabel

    return "\n".join(w.text() for w in panel.findChildren(QLabel))


def test_dialog_material_dan_panel_kebutuhan_bahan(db_sementara, qapp):
    from gui import tema
    from gui.estimasi_page import EstimasiPage
    from gui.material_dialog import MaterialDialog, MaterialElemenDialog

    tema.terapkan(qapp, "gelap")
    pid = _proyek_ac20()
    hal = EstimasiPage(pid, "AC20")

    d = MaterialDialog(pid, hal._data, hal)
    assert set(d.combo) == set(SLOT)
    cb = d.combo["DND.BATA"]
    assert cb.currentData() == "ringan10" and "(bawaan)" in cb.currentText() and "Rp" in cb.currentText()
    assert "Bata ringan" in d._info["DND.BATA"].text() and "proyek:" in d._info["DND.BATA"].text()
    assert "42.800" in d._info["DND.BATA"].toolTip()  # catatan harga mortar HSPK (rinci di tooltip)
    assert "perlu dicek" in d._info["DND.BATA"].text()
    assert not d.wadah_lanjutan.isVisibleTo(d)  # mutu beton & fondasi di bagian lanjutan
    d.btn_lanjutan.setChecked(True)
    assert d.wadah_lanjutan.isVisibleTo(d) and "BTN.KOLOM" in d.combo
    assert not d.btn_simpan.isEnabled()
    cb.setCurrentIndex(cb.findData("merah"))
    assert d.btn_simpan.isEnabled() and "naik" in d.label_selisih.text()  # bata merah lebih mahal
    d._simpan()
    assert d.diubah and _baris(pid, "DND.BATA")[0]["kode_ahsp"] == "DND.BATA~merah"

    # panel rincian per elemen: kebutuhan bahan pekerjaan + seluruh elemen
    hal.muat()
    r = next(x for x in hal._data if x["kode_ahsp"] == "DND.BATA~merah")
    sebaris = [x for x in hal._data if x["elemen_id"] == r["elemen_id"]]
    hal.panel.tampilkan_elemen(r, sebaris, hal.komponen())
    teks = _teks_panel(hal.panel)
    assert "KEBUTUHAN BAHAN" in teks and "Bata merah" in teks and "zak" in teks
    assert "ELEMEN INI" in teks

    # dialog material elemen: tabel perbandingan semua pilihan untuk volume elemen ini
    de = MaterialElemenDialog(r["elemen_id"], r["nama_elemen"], hal)
    t = de.tabel["DND.BATA"]
    assert t.rowCount() == len(SLOT["DND.BATA"].opsi)
    baris_ringan = next(i for i in range(t.rowCount()) if t.item(i, 0).data(0x0100) == "ringan10")
    assert "Bata ringan" in t.item(baris_ringan, 2).text()
    de._pilih_baris("DND.BATA", t, baris_ringan)
    assert de.combo["DND.BATA"].currentData() == "ringan10" and de.btn_simpan.isEnabled()
    de._simpan()
    kode = {x["elemen_id"]: x["kode_ahsp"] for x in _baris(pid, "DND.BATA")}
    assert kode[r["elemen_id"]] == "DND.BATA"  # elemen ini bata ringan, lainnya bata merah
    assert "DND.BATA~merah" in set(kode.values())


def test_mode_kebutuhan_bahan_keseluruhan_dan_per_lantai(db_sementara, qapp):
    from gui.estimasi_page import SEMUA_LANTAI, EstimasiPage

    pid = _proyek_ac20()
    hal = EstimasiPage(pid, "AC20")
    hal.grup_mode.button(hal.MODE_BAHAN).setChecked(True)
    hal._isi_ulang()
    nama = [hal.tabel.item(i, 1).text() for i in range(hal.tabel.rowCount()) if hal.tabel.item(i, 1)]
    assert "BAHAN / MATERIAL" in nama and "TENAGA KERJA" in nama
    assert any("Bata ringan" in x for x in nama) and any("Semen" in x for x in nama)
    semua = {k: d["jumlah"] for k, d in hal._bahan.items()}
    bata = next(d for d in hal._bahan.values() if d["nama"].startswith("Bata ringan tebal 10"))
    luas = sum(r["volume_pekerjaan"] for r in hal._data if r["kode_ahsp"] == "DND.BATA")
    assert bata["jumlah"] == pytest.approx(8.75 * luas)  # koefisien HSPK 3.6.4.2
    assert "Pasangan Dinding Bata Ringan 60x20x10 cm" in bata["pekerjaan"]
    assert len(bata["lantai"]) >= 2  # AC20: dinding di lantai dasar dan lantai atas
    # panel asal kebutuhan
    baris = next(i for i in range(hal.tabel.rowCount())
                 if hal.tabel.item(i, 1) and hal.tabel.item(i, 1).text() == bata["nama"])
    hal.tabel.selectRow(baris)
    teks = _teks_panel(hal.panel)
    assert "DARI PEKERJAAN" in teks and "PER LANTAI" in teks
    # satu lantai: kebutuhan lebih kecil dari keseluruhan
    lantai = [hal.combo_lantai.itemText(i) for i in range(hal.combo_lantai.count())
              if hal.combo_lantai.itemText(i) != SEMUA_LANTAI]
    hal.combo_lantai.setCurrentText(lantai[0])
    bata1 = next((d for d in hal._bahan.values() if d["nama"] == bata["nama"]), None)
    assert bata1 is None or bata1["jumlah"] < bata["jumlah"]
    assert sum(semua.values()) > 0


def test_export_kebutuhan_bahan_proyek_dan_per_elemen(db_sementara, tmp_path):
    from openpyxl import load_workbook

    from database.material_repository import simpan_material_elemen, simpan_pilihan
    from export_service import OpsiExport, ambil_data_export, export_excel, export_pdf

    pid = _proyek_ac20()
    simpan_pilihan(pid, {"KRM.LANTAI": "gr60"})
    eid = _baris(pid, "DND.BATA")[0]["elemen_id"]
    simpan_material_elemen(eid, {"DND.BATA": "merah"})
    data = ambil_data_export(pid)
    meta = {"nama_proyek": "AC20"}
    opsi = OpsiExport(rekap=False, rinci=False, besi=False, detail=False, lantai=False, per_lantai=False)
    path = export_excel(tmp_path / "bahan.xlsx", data, meta, opsi)
    wb = load_workbook(path)
    assert wb.sheetnames == ["RAB", "Kebutuhan Bahan", "Kebutuhan per Elemen"]
    teks = [str(c.value) for row in wb["Kebutuhan Bahan"].iter_rows() for c in row if c.value is not None]
    assert "PILIHAN MATERIAL PROYEK" in teks and "Granit 60×60 cm" in teks and "diubah" in teks
    assert "Bata merah ½ batu, 1SP : 4PP" in teks and "khusus elemen" in teks
    assert any(t.startswith("Bata ringan tebal 10") for t in teks) and "Bata merah" in teks
    per_el = [str(c.value) for row in wb["Kebutuhan per Elemen"].iter_rows() for c in row if c.value is not None]
    assert any("– Bata merah" in t for t in per_el) and any("– Bata ringan" in t for t in per_el)
    assert any("Pasangan Dinding Bata Merah" in t for t in per_el)
    pdf = export_pdf(tmp_path / "bahan.pdf", data, meta, opsi)
    assert (tmp_path / "bahan.pdf").stat().st_size > 2000 and pdf.endswith(".pdf")


def test_item_tambahan_dari_katalog(db_sementara, tmp_path):
    from database.berkas_proyek import duplikat_proyek
    from database.estimasi_repository import NilaiTidakValid, get_harga_satuan_pekerjaan, get_hasil_estimasi_by_proyek
    from database.katalog import cari_katalog, daftar_katalog, hapus_item, tambah_item
    from estimasi_service import hitung_ulang_dari_elemen, jalankan_estimasi

    data = daftar_katalog()
    assert len(data) == len(AHSP)
    kloset = cari_katalog(data, "closet duduk")
    assert kloset and kloset[0]["kode"] == "3.18.3.1" and kloset[0]["kategori"] == "Sanitair"
    assert all(d["bab"] == "5" for d in cari_katalog(data, "", "5"))

    pid = _proyek_ac20()
    lantai = get_hasil_estimasi_by_proyek(pid)[0]["lantai"]
    hid = tambah_item(pid, "3.18.3.1", 2, lantai, "KM utama & KM tamu")
    hid2 = tambah_item(pid, "5.1.1.1" if "5.1.1.1" in AHSP else next(k for k in AHSP if k.startswith("5.1.")), 10)
    rows = {r["hasil_id"]: r for r in get_hasil_estimasi_by_proyek(pid)}
    r = rows[hid]
    assert r["kode_ahsp"] == "HSPK-3.18.3.1" and r["item_tambahan"] == 1 and r["elemen_id"] is None
    assert r["lantai"] == lantai and r["catatan"] == "KM utama & KM tamu"
    # harga = komponen analisa x (1 + BUK). (Untuk 3.18.3.1 harga F dokumen sendiri tidak sama dengan
    # hitungan komponennya; selisih itu dilaporkan tools/impor_hspk.py.)
    dasar = sum(k[3] * k[4] for k in AHSP["3.18.3.1"][3])
    assert get_harga_satuan_pekerjaan(r["pekerjaan_id"]) == pytest.approx(dasar * 1.1, rel=1e-6)
    assert r["subtotal_biaya"] == pytest.approx(2 * get_harga_satuan_pekerjaan(r["pekerjaan_id"]))
    assert rows[hid2]["lantai"] is None

    # estimasi ulang (dari elemen maupun IFC) tidak menghapus item tambahan
    hitung_ulang_dari_elemen(pid)
    jalankan_estimasi(pid)
    assert {hid, hid2} <= {r["hasil_id"] for r in get_hasil_estimasi_by_proyek(pid)}
    # ikut duplikat proyek
    salinan = duplikat_proyek(pid, "Salinan")
    assert any(r["kode_ahsp"] == "HSPK-3.18.3.1" and r["item_tambahan"] for r in get_hasil_estimasi_by_proyek(salinan))

    with pytest.raises(NilaiTidakValid):
        tambah_item(pid, "3.18.3.1", 0)
    with pytest.raises(NilaiTidakValid):
        hapus_item(next(r["hasil_id"] for r in get_hasil_estimasi_by_proyek(pid) if not r["item_tambahan"]))
    hapus_item(hid)
    assert hid not in {r["hasil_id"] for r in get_hasil_estimasi_by_proyek(pid)}


def test_dialog_katalog_tambah_dan_hapus_item(db_sementara, qapp, monkeypatch):
    from gui import tema
    from gui.estimasi_page import EstimasiPage
    from gui.katalog_dialog import KatalogDialog

    pid = _proyek_ac20()
    hal = EstimasiPage(pid, "AC20")
    lantai = sorted({r["lantai"] for r in hal._data if r["lantai"]})
    d = KatalogDialog(pid, lantai, hal)
    assert d.tabel.rowCount() > 100  # dibatasi MAKS_BARIS
    d.kolom_cari.setText("closet duduk")
    d._isi()
    assert d.tabel.rowCount() >= 1 and d.tabel.item(0, 0).text() == "3.18.3.1"
    d.tabel.selectRow(0)
    assert d.btn_tambah.isEnabled() and "Bahan" in d.label_analisa.text() and d.label_satuan.text() == "buah"
    d.spin_volume.setValue(2)
    d.combo_lantai.setCurrentIndex(1)
    d._tambah()
    d.combo_bab.setCurrentIndex(d.combo_bab.findData("5"))
    d.kolom_cari.setText("")
    d._isi()
    assert all(d.tabel.item(i, 0).text().startswith("5.") for i in range(min(20, d.tabel.rowCount())))
    assert len(d.ditambah) == 1

    hal.muat()
    r = next(x for x in hal._data if x["item_tambahan"])
    assert r["kategori"] == "Sanitair" and r["lantai"] == lantai[0] and r["volume_pekerjaan"] == 2
    for mode in range(5):  # semua mode tampil tanpa galat
        hal.grup_mode.button(mode).setChecked(True)
        hal._isi_ulang()
    hal.grup_mode.button(hal.MODE_DETAIL).setChecked(True)
    hal._isi_ulang()
    hal.panel.tampilkan_elemen(r, [r], hal.komponen())
    assert "Hapus Item Tambahan" in [b.text() for b in hal.panel.findChildren(type(d.btn_tambah))]
    monkeypatch.setattr(tema, "tanya", lambda *a, **k: True)
    hal._hapus_item(r["hasil_id"])
    assert not any(x["item_tambahan"] for x in hal._data)
    hal.riwayat.undo()  # urungkan hapus
    assert any(x["item_tambahan"] for x in hal._data) or any(
        x["item_tambahan"] for x in __import__("database.estimasi_repository", fromlist=["x"]).get_hasil_estimasi_by_proyek(pid))


def test_ganti_bahan_di_tabel_per_elemen_mengubah_total_dan_kebutuhan(db_sementara, qapp, monkeypatch):
    """Ganti bahan satu dinding di mode Per Elemen: subtotal baris, total RAB (kartu), rekap per lantai,
    dan kebutuhan bahan ikut berubah; urungkan mengembalikan semuanya."""
    from gui import tema
    from gui.estimasi_page import EstimasiPage, _bahan_baris
    from database.biaya_repository import ringkasan_biaya

    pesan = []
    monkeypatch.setattr(tema, "toast", lambda w, teks, *a, **k: pesan.append(teks))
    pid = _proyek_ac20()
    hal = EstimasiPage(pid, "AC20")
    hal.grup_mode.button(hal.MODE_DETAIL).setChecked(True)
    hal._isi_ulang()
    t = hal.tabel
    i = next(i for i in range(t.rowCount()) if t.item(i, 0).data(0x0100)[0] == "detail"
             and hal._index[t.item(i, 0).data(0x0100)[1]]["kode_ahsp"] == "DND.BATA")
    r = hal._index[t.item(i, 0).data(0x0100)[1]]
    cb = t.cellWidget(i, 2)
    assert cb is not None and cb.currentData() == "ringan10" and _bahan_baris(r) == ("DND.BATA", "ringan10")
    total_lama = ringkasan_biaya(pid)
    lantai_lama = {x["lantai"]: x["total"] for x in __import__("rab_rinci").rekap_per_lantai(hal._data)}
    kartu_lama = hal.k_total.findChildren(type(hal.label_mode))[1].text() if hal.k_total.findChildren(type(hal.label_mode)) else ""

    cb.setCurrentIndex(cb.findData("merah"))  # langsung tersimpan

    baru = {x["hasil_id"]: x for x in hal._data}[r["hasil_id"]]
    assert baru["kode_ahsp"] == "DND.BATA~merah"
    assert baru["volume_pekerjaan"] == pytest.approx(r["volume_pekerjaan"])  # volume tetap
    assert baru["subtotal_biaya"] > r["subtotal_biaya"]  # bata merah lebih mahal
    d = baru["subtotal_biaya"] - r["subtotal_biaya"]
    assert ringkasan_biaya(pid)["langsung"] == pytest.approx(total_lama["langsung"] + d)
    lantai_baru = {x["lantai"]: x["total"] for x in __import__("rab_rinci").rekap_per_lantai(hal._data)}
    assert lantai_baru[r["lantai"]] == pytest.approx(lantai_lama[r["lantai"]] + d)
    assert pesan and "Total RAB" in pesan[-1] and "naik" in pesan[-1]
    if kartu_lama:
        assert hal.k_total.findChildren(type(hal.label_mode))[1].text() != kartu_lama
    # kebutuhan bahan proyek kini memuat bata merah, bata ringan berkurang
    hal.grup_mode.button(hal.MODE_BAHAN).setChecked(True)
    hal._isi_ulang()
    nama = {x["nama"]: x["jumlah"] for x in hal._bahan.values()}
    assert any(n.startswith("Bata merah") for n in nama)
    # urungkan: kembali bata ringan dan total semula
    hal.riwayat.undo()
    assert ringkasan_biaya(pid)["langsung"] == pytest.approx(total_lama["langsung"])



def test_per_elemen_dikelompokkan_per_lantai_dan_jenis(db_sementara, qapp):
    from gui.estimasi_page import EstimasiPage, kelompok_lantai_jenis
    from rab_rinci import daftar_lantai

    pid = _proyek_ac20()
    hal = EstimasiPage(pid, "AC20")
    kel = kelompok_lantai_jenis(hal._data)
    assert [g["lantai"] for g in kel] == daftar_lantai(hal._data)  # urut elevasi, sama dengan Per Lantai
    assert sum(g["total"] for g in kel) == pytest.approx(sum(r["subtotal_biaya"] for r in hal._data))
    dasar = kel[0]
    jenis = [j["jenis"] for j in dasar["jenis"]]
    assert "Dinding" in jenis
    dinding = next(j for j in dasar["jenis"] if j["jenis"] == "Dinding")
    assert dinding["jumlah_elemen"] == len({r["elemen_id"] for r in hal._data
                                           if r["lantai"] == dasar["lantai"] and r["kelas"] == "WALL"})

    hal.grup_mode.button(hal.MODE_DETAIL).setChecked(True)
    hal._isi_ulang()
    t = hal.tabel
    data = [t.item(i, 0).data(0x0100) for i in range(t.rowCount())]
    assert data[0] == ("lantai_el", dasar["lantai"]) and data[1][0] == "jenis_el"
    assert sum(1 for d in data if d[0] == "detail") == len(hal._data)
    # baris judul jenis -> panel ringkasan kelompok: jumlah elemen, volume pekerjaan, kebutuhan bahan
    i = data.index(("jenis_el", (dasar["lantai"], "Dinding")))
    t.selectRow(i)
    teks = _teks_panel(hal.panel)
    assert "RINGKASAN KELOMPOK" in teks and "Dinding" in teks and str(dinding["jumlah_elemen"]) in teks
    assert "VOLUME PEKERJAAN" in teks and "KEBUTUHAN BAHAN" in teks and "Bata ringan" in teks
    t.selectRow(0)
    assert "PER JENIS ELEMEN" in _teks_panel(hal.panel)
