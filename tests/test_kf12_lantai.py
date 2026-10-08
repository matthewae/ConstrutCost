"""KF-12 Multi lantai: rekap biaya per lantai, filter lantai, dan export."""

import os

import pytest

from conftest import AC20, DUPLEX


@pytest.fixture
def baris_ac20(db_sementara):
    from database.estimasi_repository import get_hasil_estimasi_by_proyek
    from database.proyek_repository import create_proyek
    from estimasi_service import jalankan_estimasi

    pid = create_proyek("AC20", str(AC20))
    jalankan_estimasi(pid)
    return pid, [dict(r) for r in get_hasil_estimasi_by_proyek(pid)]


def test_rekap_per_lantai_urut_elevasi_dan_total_sama(baris_ac20):
    from rab_rinci import daftar_lantai, rekap_per_lantai

    _, baris = baris_ac20
    assert daftar_lantai(baris) == ["Erdgeschoss", "Dachgeschoss"]  # elevasi 0 lalu 2,7 m
    rekap = rekap_per_lantai(baris)
    assert [x["lantai"] for x in rekap] == ["Erdgeschoss", "Dachgeschoss"]
    assert sum(x["total"] for x in rekap) == pytest.approx(sum(r["subtotal_biaya"] for r in baris))
    eg = rekap[0]
    # pekerjaan tanah & fondasi turunan hanya ada di lantai dasar
    assert {"Tanah", "Fondasi"} <= {k["kategori"] for k in eg["kategori"]}
    assert "Tanah" not in {k["kategori"] for k in rekap[1]["kategori"]}
    assert eg["jumlah_elemen"] == len({r["elemen_id"] for r in baris if r["lantai"] == "Erdgeschoss"})


def test_baris_tanpa_lantai(baris_ac20):
    from rab_rinci import TANPA_LANTAI, rekap_per_lantai

    _, baris = baris_ac20
    baris[0] = {**baris[0], "lantai": None, "elevasi_lantai": None}
    assert rekap_per_lantai(baris)[-1]["lantai"] == TANPA_LANTAI


def test_filter_lantai_di_halaman(baris_ac20):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication

    app = QApplication.instance() or QApplication([])
    from gui import tema
    from gui.estimasi_page import EstimasiPage

    tema.terapkan(app, "gelap")
    pid, baris = baris_ac20
    hal = EstimasiPage(pid, "AC20")
    assert [hal.combo_lantai.itemText(i) for i in range(hal.combo_lantai.count())] == [
        "Semua lantai", "Erdgeschoss", "Dachgeschoss",
    ]
    hal.grup_mode.button(hal.MODE_LANTAI).click()
    assert hal.label_jumlah.text().startswith("2 lantai")
    hal._buka_lantai("Dachgeschoss")  # klik dua kali nama lantai: rincian kebutuhan lantai itu
    assert hal.mode == hal.MODE_LANTAI and hal.combo_lantai.currentText() == "Dachgeschoss"
    total_dg = sum(r["subtotal_biaya"] for r in baris if r["lantai"] == "Dachgeschoss")
    assert tema.format_rupiah(total_dg) in hal.label_jumlah.text()


def test_export_rekap_per_lantai(db_sementara, tmp_path):
    from openpyxl import load_workbook

    from database.proyek_repository import create_proyek
    from estimasi_service import jalankan_estimasi
    from export_service import OpsiExport, ambil_data_export, export_excel, export_pdf

    pid = create_proyek("Duplex", str(DUPLEX))
    jalankan_estimasi(pid)
    data = ambil_data_export(pid)
    path = tmp_path / "rab.xlsx"
    export_excel(path, data, {}, OpsiExport())
    ws = load_workbook(path)["Rekap per Lantai"]
    teks = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any(t.startswith("Level 1") for t in teks)
    assert "JUMLAH 4 LANTAI" in teks
    export_pdf(tmp_path / "rab.pdf", data, {}, OpsiExport(detail=False))
    from pypdf import PdfReader

    isi = " ".join(p.extract_text() for p in PdfReader(tmp_path / "rab.pdf").pages)
    assert "REKAPITULASI BIAYA & KEBUTUHAN STRUKTUR PER LANTAI" in isi
