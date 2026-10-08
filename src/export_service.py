"""
Export hasil estimasi RAB ke Excel (.xlsx) dan PDF (KF-8, KF-13, KF-17).

Isi dokumen:
- RAB              : bagian per kategori (I, II, ...), item per pekerjaan, jumlah tiap bagian, lalu
                     biaya langsung (A), biaya tidak langsung (B), PPN, total, dan pembulatan.
- Rekapitulasi     : ringkasan biaya (KF-13): A per kategori, rincian B, PPN, total, terbilang.
- RAB Rinci        : beton, bekisting, dan tulangan per tipe elemen (Kolom K1 20/25: 6 D13, ...).
- Kebutuhan Besi   : berat, panjang, dan jumlah batang 12 m per diameter.
- Rekap per Lantai : biaya, beton, bekisting, dan besi tiap lantai dirinci per kategori (KF-12).
- Kebutuhan per Lantai : tabel silang besi per diameter, beton per mutu, bahan, tenaga kerja, dan alat
                     untuk setiap lantai (kolom lantai sebanyak lantai di model).
- Rincian Perhitungan Lantai : berurutan dari lantai terbawah sampai penutup bangunan (atap / dak /
                     plafon): volume pekerjaan, beton, bekisting, besi per diameter (berat ÷ berat/m' =
                     panjang -> batang 12 m), dan bahan per lantai.
- Lt01 ... LtNN    : satu sheet / bagian per lantai: biaya per kategori, struktur beton per tipe, besi per
                     diameter, bahan, tenaga kerja, alat (koefisien AHSP), dan RAB rinci lantai.
- Detail Elemen    : volume tiap elemen dan lantai (KF-12), opsional dengan uraian rumus (KF-18).

KF-17: kolom No, Kode Analisa, Harga Satuan, Bobot (%), dan Rumus dapat dipilih; Uraian, Volume,
Satuan, dan Jumlah Harga selalu ada. Excel memakai rumus aktif (jumlah = volume x harga satuan).
"""

import logging
import os
import re
import tempfile
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from aktivitas import catat
from database.biaya_repository import ringkasan_biaya
from database.estimasi_repository import BUK_RATE, PPN_RATE, _connect, get_hasil_estimasi_by_proyek
from database.penulangan_repository import daftar_tipe, kebutuhan_besi
from kebutuhan_lantai import (
    BERAT_ZAK_SEMEN,
    TIPE_SUMBER_DAYA,
    komponen_pekerjaan,
    penutup_bangunan,
    rab_lantai,
    matriks_lantai,
    rasio_wajar,
    rincian_per_lantai,
    ringkas_lantai,
)
from rab_rinci import URUTAN_KATEGORI, susun_rinci, urutan_kategori

log = logging.getLogger("coststruct.export")

__all__ = ["URUTAN_KATEGORI", "OpsiExport", "ambil_data_export", "export_excel", "export_pdf", "kelompokkan"]

# KF-17: kolom yang bisa dipilih pengguna (kunci -> label)
KOLOM_OPSIONAL = {
    "no": "No",
    "kode": "Kode Analisa",
    "harga": "Harga Satuan",
    "bobot": "Bobot (%)",
    "rumus": "Uraian Rumus (detail elemen)",
}
KOLOM_BAWAAN = ("no", "kode", "harga")

# Beda kedua bentuk RAB, dicetak di bawah judul sheet / halaman agar pembaca laporan tidak bingung.
KET_RAB = ("Rekap RAB: satu baris untuk setiap item pekerjaan; volume = jumlah dari semua elemen di model "
           "(bentuk ringkas untuk dokumen RAB / penawaran). Rincian per tipe elemen ada di sheet RAB Rinci.")
KET_RAB_RINCI = ("RAB Rinci: setiap item pekerjaan dipecah. Struktur per tipe elemen seperti RAP konsultan (mis. Kolom K1 "
                 "(20/25) — 12 buah: beton, bekisting, tulangan utama, sengkang); pekerjaan lain per tipe elemen IFC (mis. "
                 "Pasangan dinding bata: Exterior Brick — 16 buah, Interior 100 mm — 9 buah). Total biayanya sama dengan sheet RAB.")


@dataclass
class OpsiExport:
    kolom: tuple = KOLOM_BAWAAN
    rekap: bool = True
    rinci: bool = True
    besi: bool = True
    detail: bool = True
    lantai: bool = True  # KF-12 rekap per lantai
    per_lantai: bool = True  # rincian kebutuhan tiap lantai: satu sheet Excel / satu bagian PDF per lantai
    orientasi_pdf: str = "portrait"  # portrait / landscape

    def ada(self, kunci: str) -> bool:
        return kunci in self.kolom


# ---------------------------------------------------------------- data


def ambil_data_export(proyek_id: int) -> dict:
    conn = _connect()
    try:
        pr = conn.execute("SELECT nama_proyek, path_file_ifc FROM proyek WHERE id = ?", (proyek_id,)).fetchone()
    finally:
        conn.close()
    if pr is None:
        raise ValueError("Proyek tidak ditemukan.")
    baris = [dict(r) for r in get_hasil_estimasi_by_proyek(proyek_id)]
    if not baris:
        raise ValueError("Belum ada hasil estimasi untuk di-export.")
    return {
        "proyek_id": proyek_id,
        "nama_proyek": pr["nama_proyek"],
        "path_ifc": pr["path_file_ifc"],
        "baris": baris,
        "ringkasan": ringkasan_biaya(proyek_id),
        "besi": kebutuhan_besi(proyek_id),
        "tipe": daftar_tipe(proyek_id),
        "komponen": komponen_pekerjaan(r["pekerjaan_id"] for r in baris),  # analisa AHSP: kebutuhan bahan
    }


def kelompokkan(baris: list) -> list:
    """Gabungkan semua elemen per pekerjaan, lalu kelompokkan per kategori.
    Return: [{'kategori', 'items': [{'kode','nama','satuan','volume','harga','jumlah'}], 'total'}]"""
    per_pekerjaan = {}
    for r in baris:
        d = per_pekerjaan.setdefault(r["pekerjaan_id"], {
            "kategori": r["kategori"], "kode": r["kode_ahsp"] or "", "nama": r["nama_pekerjaan"],
            "satuan": r["satuan"], "volume": 0.0, "jumlah": 0.0,
        })
        d["volume"] += r["volume_pekerjaan"]
        d["jumlah"] += r["subtotal_biaya"]
    per_kategori = {}
    for d in per_pekerjaan.values():
        d["harga"] = d["jumlah"] / d["volume"] if d["volume"] else 0.0
        per_kategori.setdefault(d["kategori"], []).append(d)
    hasil = []
    for k in sorted(per_kategori, key=urutan_kategori):
        items = sorted(per_kategori[k], key=lambda x: x["nama"])
        hasil.append({"kategori": k, "items": items, "total": sum(i["jumlah"] for i in items)})
    return hasil


def _entri_rab(baris: list) -> list:
    """Baris tabel RAB per pekerjaan: kategori, item, subtotal."""
    out = []
    for i, k in enumerate(kelompokkan(baris), 1):
        out.append({"jenis": "kategori", "no": f"{_romawi(i)}.", "uraian": k["kategori"].upper()})
        for n, it in enumerate(k["items"], 1):
            out.append({"jenis": "item", "no": str(n), "kode": it["kode"], "uraian": it["nama"],
                        "volume": it["volume"], "satuan": it["satuan"], "harga": it["harga"], "jumlah": it["jumlah"]})
        out.append({"jenis": "subtotal", "uraian": f"Jumlah {_romawi(i)}. {k['kategori']}", "jumlah": k["total"]})
    return out


def _entri_rinci(baris: list, per_tipe_elemen: bool = True) -> list:
    """Baris tabel RAB rinci: struktur per tipe penulangan, pekerjaan lain per tipe elemen IFC."""
    out = []
    for i, k in enumerate(susun_rinci(baris, per_tipe_elemen), 1):
        out.append({"jenis": "kategori", "no": f"{_romawi(i)}.", "uraian": k["kategori"].upper()})
        nomor = 0
        for g in k["grup"]:
            if g["judul"]:
                nomor += 1
                out.append({"jenis": "grup", "no": str(nomor), "uraian": g["judul"]})
            for it in g["items"]:
                if g["judul"]:
                    no = "-"
                else:
                    nomor += 1
                    no = str(nomor)
                out.append({"jenis": "item", "no": no, "kode": it["kode"], "uraian": it["label"],
                            "volume": it["volume"], "satuan": it["satuan"], "harga": it["harga"],
                            "jumlah": it["jumlah"], "indent": bool(g["judul"])})
        out.append({"jenis": "subtotal", "uraian": f"Jumlah {_romawi(i)}. {k['kategori']}", "jumlah": k["total"]})
    return out


def baris_detail(baris: list) -> list:
    """Baris per elemen (KF-12/KF-18), diurutkan per lantai -> elemen -> kategori."""
    hasil = []
    for r in baris:
        v = r["volume_pekerjaan"]
        pekerjaan = r["nama_pekerjaan"]
        if r.get("uraian") and r.get("diameter"):
            pekerjaan += f" — {r['uraian']}"
        hasil.append({
            "lantai": r["lantai"] or "-",
            "elemen": r["nama_elemen"] or "-",
            "kategori": r["kategori"],
            "kode": r["kode_ahsp"] or "",
            "pekerjaan": pekerjaan,
            "satuan": r["satuan"],
            "volume": v,
            "harga": r["subtotal_biaya"] / v if v else 0.0,
            "jumlah": r["subtotal_biaya"],
            "status": "Manual" if r["diedit_manual"] else ("Dimensi diubah" if r.get("dimensi_manual") else "Otomatis"),
            "rumus": r.get("rumus") or "",
            "catatan": r.get("catatan") or "",
        })
    return sorted(hasil, key=lambda x: (x["lantai"], x["elemen"], x["kategori"], x["pekerjaan"]))


def _rincian_lantai(data: dict) -> list:
    """Rincian per lantai (kebutuhan_lantai), dihitung sekali per export."""
    if "_rincian_lantai" not in data:
        data["_rincian_lantai"] = rincian_per_lantai(data["baris"], data.get("komponen"))
    return data["_rincian_lantai"]


def nama_sheet_lantai(i: int, lantai: str, terpakai: set) -> str:
    """Nama sheet Excel untuk lantai ke-i: 'Lt01 LANTAI 1'. Maks. 31 karakter, tanpa karakter terlarang
    ([]:*?/\\), dan unik walau nama lantai di model sama atau sangat panjang."""
    bersih = re.sub(r"[\[\]:*?/\\]", " ", lantai)
    bersih = " ".join(bersih.split()).strip("'") or "Lantai"
    nama = f"Lt{i:02d} {bersih}"[:31].rstrip()
    dasar, n = nama, 2
    while nama.lower() in terpakai:
        akhir = f" ({n})"
        nama = dasar[:31 - len(akhir)].rstrip() + akhir
        n += 1
    terpakai.add(nama.lower())
    return nama


class FolderTidakBisaDitulis(PermissionError):
    """Folder tujuan menolak file baru (izin folder, Controlled Folder Access Windows, drive read-only)."""


def tulis_aman(path, tulis) -> Path:
    """Tulis file export lewat file sementara di folder tujuan, lalu ganti file lama sekaligus.

    - Folder tidak bisa ditulis -> FolderTidakBisaDitulis (pesan & saran khusus, KF-14).
    - File tujuan sedang dikunci program lain (Excel, panel Preview File Explorer, antivirus / OneDrive
      yang sedang memindai) -> hasil disimpan dengan nama baru "nama (2).xlsx" dan path itu dikembalikan,
      sehingga export tidak gagal hanya karena file lama masih terbuka.
    - Gagal di tengah jalan -> file lama tidak rusak, file sementara dihapus."""
    path = Path(path)
    try:
        fd, tmp = tempfile.mkstemp(prefix="~coststruct_", suffix=path.suffix, dir=path.parent)
    except PermissionError as e:
        raise FolderTidakBisaDitulis(
            e.errno, f"Folder tujuan tidak mengizinkan file baru dibuat: {path.parent}", str(path.parent)
        ) from e
    os.close(fd)
    try:
        tulis(tmp)
        for n in range(1, 50):
            tujuan = path if n == 1 else path.with_name(f"{path.stem} ({n}){path.suffix}")
            try:
                if n == 1:  # antivirus / OneDrive sering mengunci sesaat: coba ulang dulu sebelum ganti nama
                    for _ in range(3):
                        try:
                            os.replace(tmp, tujuan)
                            return tujuan
                        except PermissionError:
                            time.sleep(0.2)
                os.replace(tmp, tujuan)
                if n > 1:
                    log.warning("%s sedang dikunci program lain; hasil disimpan sebagai %s", path.name, tujuan.name)
                return tujuan
            except PermissionError:
                continue  # file dengan nama itu sedang terbuka / terkunci: coba nama berikutnya
        raise PermissionError(13, "Semua nama file alternatif sedang dikunci", str(path))
    finally:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass


def nama_file_default(nama_proyek: str, ekstensi: str) -> str:
    slug = re.sub(r"[^\w\-]+", "_", nama_proyek, flags=re.UNICODE).strip("_") or "Proyek"
    return f"RAB_{slug}_{datetime.now():%Y%m%d}.{ekstensi}"


def _romawi(n: int) -> str:
    out = ""
    for nilai, sim in ((10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")):
        while n >= nilai:
            out += sim
            n -= nilai
    return out


def _rp(x: float, desimal: int = 2) -> str:
    """Format Indonesia: 1.234.567,89"""
    return f"{x:,.{desimal}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _persen(x: float) -> str:
    return f"{x * 100:.2f}".replace(".", ",") + "%"


def _info_dokumen(data: dict, meta: dict) -> list:
    info = [("Pekerjaan", meta.get("nama_proyek") or data["nama_proyek"])]
    if meta.get("lokasi"):
        info.append(("Lokasi", meta["lokasi"]))
    if meta.get("pemilik"):
        info.append(("Pemilik / Instansi", meta["pemilik"]))
    if meta.get("tahun"):
        info.append(("Tahun Anggaran", str(meta["tahun"])))
    if data.get("path_ifc"):
        info.append(("Sumber QTO", f"Model IFC: {Path(data['path_ifc']).name}"))
    return info


def _label_btl(b: dict) -> str:
    if b["jenis"] == "persen":
        return f"{b['uraian']} ({_rp(b['nilai'], 2).rstrip('0').rstrip(',')}% × A)"
    return b["uraian"]


CATATAN = [
    "Volume dihitung otomatis dari model IFC (quantity take-off) dan dapat diedit manual pada aplikasi.",
    f"Harga satuan = analisa harga satuan (bahan + upah + alat) ditambah Biaya Umum & Keuntungan {BUK_RATE:.0%}.",
    "Pembesian dihitung per diameter dari tipe penulangan: W = Σ(π d²/4 × 7850 × L_eff × n) × 1,05.",
    f"PPN {PPN_RATE:.0%} dihitung dari jumlah biaya langsung dan biaya tidak langsung.",
]


def _kolom_tabel(opsi: OpsiExport, jenis: str = "rab", ada_catatan: bool = False) -> list:
    """Daftar (kunci, judul, lebar_excel, lebar_pdf_relatif) sesuai pilihan kolom (KF-17)."""
    k = []
    if opsi.ada("no"):
        k.append(("no", "No", 6, 0.6))
    if jenis == "detail":
        k += [("lantai", "Lantai", 14, 1.3), ("elemen", "Elemen", 26, 2.2)]
    if opsi.ada("kode"):
        k.append(("kode", "Kode Analisa", 16, 1.75))
    k.append(("uraian", "Uraian Pekerjaan", 52 if jenis != "detail" else 40, 4.6 if jenis != "detail" else 3.4))
    k += [("volume", "Volume", 13, 1.3), ("satuan", "Sat", 7, 0.75)]
    if opsi.ada("harga"):
        k.append(("harga", "Harga Satuan (Rp)", 18, 1.6))
    k.append(("jumlah", "Jumlah Harga (Rp)", 20, 1.8))
    if opsi.ada("bobot") and jenis != "detail":
        k.append(("bobot", "Bobot", 10, 1.0))
    if jenis == "detail":
        k.append(("status", "Status", 12, 1.1))
        if ada_catatan:
            k.append(("catatan", "Catatan", 30, 2.4))
        if opsi.ada("rumus"):
            k.append(("rumus", "Uraian Rumus", 70, 5.0))
    return k


# ---------------------------------------------------------------- Excel


def export_excel(path, data: dict, meta: dict, opsi: OpsiExport | None = None) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L

    opsi = opsi or OpsiExport()
    f_norm = Font(name="Arial", size=10)
    f_bold = Font(name="Arial", size=10, bold=True)
    f_judul = Font(name="Arial", size=14, bold=True)
    f_input = Font(name="Arial", size=10, color="0000FF")
    f_catatan = Font(name="Arial", size=9, italic=True)
    fill_head = PatternFill("solid", fgColor="D9D9D9")
    fill_kat = PatternFill("solid", fgColor="EEF3F8")
    fill_grup = PatternFill("solid", fgColor="F7F9FB")
    sisi = Side(style="thin", color="808080")
    kotak = Border(left=sisi, right=sisi, top=sisi, bottom=sisi)
    ANGKA, ANGKA3, PERSEN = "#,##0.00", "#,##0.000", "0.00%"
    ring = data["ringkasan"]
    wb = Workbook()

    def judul_sheet(ws, teks, n_kolom):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_kolom)
        ws["A1"] = teks
        ws["A1"].font = f_judul
        ws["A1"].alignment = Alignment(horizontal="center")
        r = 2
        for label, nilai in _info_dokumen(data, meta):
            ws.cell(r, 1, f"{label} : {nilai}").font = f_norm
            r += 1
        return r + 1

    def header(ws, r, judul):
        for c, teks in enumerate(judul, 1):
            s = ws.cell(r, c, teks)
            s.font, s.fill, s.border = f_bold, fill_head, kotak
            s.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def gaya_baris(ws, r, n, font=f_norm, fill=None):
        for c in range(1, n + 1):
            s = ws.cell(r, c)
            s.font, s.border = font, kotak
            if fill:
                s.fill = fill

    def atur_cetak(ws, baris_header, orientasi="portrait"):
        ws.page_setup.orientation = orientasi
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f"{baris_header}:{baris_header}"
        ws.freeze_panes = ws.cell(baris_header + 1, 1)
        ws.page_margins.left = ws.page_margins.right = 0.5

    def tulis_tabel(ws, r, kolom, entri):
        """Tulis tabel RAB (kategori / grup / item / subtotal). Return (r, baris_subtotal, baris_item)."""
        ci = {k[0]: i + 1 for i, k in enumerate(kolom)}
        n = len(kolom)
        c_uraian, c_jumlah = ci["uraian"], ci["jumlah"]
        sub, item_rows, awal = [], [], None
        for e in entri:
            if e["jenis"] == "kategori":
                gaya_baris(ws, r, n, f_bold, fill_kat)
                if "no" in ci:
                    ws.cell(r, ci["no"], e["no"])
                    ws.cell(r, c_uraian, e["uraian"])
                else:
                    ws.cell(r, c_uraian, f"{e['no']} {e['uraian']}")
                awal = r + 1
            elif e["jenis"] == "grup":
                gaya_baris(ws, r, n, f_bold, fill_grup)
                if "no" in ci:
                    ws.cell(r, ci["no"], e["no"]).alignment = Alignment(horizontal="center")
                ws.cell(r, c_uraian, e["uraian"])
            elif e["jenis"] == "item":
                gaya_baris(ws, r, n)
                for kunci in ("no", "kode", "satuan"):
                    if kunci in ci:
                        ws.cell(r, ci[kunci], e[kunci]).alignment = Alignment(
                            horizontal="center" if kunci != "kode" else "left", vertical="top"
                        )
                s = ws.cell(r, c_uraian, ("    " if e.get("indent") else "") + e["uraian"])
                s.alignment = Alignment(wrap_text=True, vertical="top")
                ws.cell(r, ci["volume"], e["volume"]).number_format = ANGKA
                if "harga" in ci:
                    ws.cell(r, ci["harga"], e["harga"]).number_format = ANGKA
                    ws.cell(r, c_jumlah, f"={L(ci['volume'])}{r}*{L(ci['harga'])}{r}")
                else:
                    ws.cell(r, c_jumlah, e["jumlah"])
                ws.cell(r, c_jumlah).number_format = ANGKA
                item_rows.append(r)
            elif e["jenis"] == "subtotal":
                gaya_baris(ws, r, n, f_bold)
                ws.cell(r, c_uraian, e["uraian"]).alignment = Alignment(horizontal="right")
                ws.cell(r, c_jumlah, f"=SUM({L(c_jumlah)}{awal}:{L(c_jumlah)}{r - 1})").number_format = ANGKA
                sub.append(r)
            r += 1
        return r, sub, item_rows

    def isi_bobot(ws, kolom, rows, ref_total):
        ci = {k[0]: i + 1 for i, k in enumerate(kolom)}
        if "bobot" not in ci:
            return
        for r in rows:
            s = ws.cell(r, ci["bobot"], f"=IF({ref_total}=0,0,{L(ci['jumlah'])}{r}/{ref_total})")
            s.number_format = PERSEN

    # ---------- Sheet RAB ----------
    ws = wb.active
    ws.title = "RAB"
    kolom = _kolom_tabel(opsi)
    for i, k in enumerate(kolom, 1):
        ws.column_dimensions[L(i)].width = k[2]
    r = judul_sheet(ws, "RENCANA ANGGARAN BIAYA (RAB)", len(kolom))
    ws.cell(r - 1, 1, KET_RAB).font = f_catatan
    hdr = r
    header(ws, r, [k[1] for k in kolom])
    r, sub_rab, item_rab = tulis_tabel(ws, r + 1, kolom, _entri_rab(data["baris"]))
    ci = {k[0]: i + 1 for i, k in enumerate(kolom)}
    cu, cj = ci["uraian"], L(ci["jumlah"])
    r += 1

    def baris_total(ws, r, teks, rumus, fmt=ANGKA, n=len(kolom)):
        gaya_baris(ws, r, n, f_bold)
        ws.cell(r, cu, teks).alignment = Alignment(horizontal="right")
        ws.cell(r, ci["jumlah"], rumus).number_format = fmt
        return r

    rA = baris_total(ws, r, "A. JUMLAH BIAYA LANGSUNG", "=" + "+".join(f"{cj}{x}" for x in sub_rab))
    baris_A = rA
    pct = sum(b["nilai"] for b in ring["item_tidak_langsung"] if b["jenis"] == "persen") / 100
    tetap = sum(b["nilai"] for b in ring["item_tidak_langsung"] if b["jenis"] == "nilai")
    rB = baris_total(ws, r + 1, "B. BIAYA TIDAK LANGSUNG (rincian di sheet Rekapitulasi)",
                     f"={cj}{rA}*{pct!r}+{tetap!r}")
    rJ = baris_total(ws, r + 2, "JUMLAH (A + B)", f"={cj}{rA}+{cj}{rB}")
    rP = baris_total(ws, r + 3, f"PPN {PPN_RATE:.0%}", f"={cj}{rJ}*{PPN_RATE!r}")
    rT = baris_total(ws, r + 4, "TOTAL", f"={cj}{rJ}+{cj}{rP}")
    baris_total(ws, r + 5, "TOTAL DIBULATKAN", f"=ROUNDDOWN({cj}{rT},-3)")
    isi_bobot(ws, kolom, item_rab + sub_rab, f"${cj}${rA}")
    r += 7
    ws.cell(r, 1, f"Terbilang: {ring['terbilang']}").font = Font(name="Arial", size=10, italic=True, bold=True)
    r += 2
    ws.cell(r, 1, "Catatan:").font = f_catatan
    for teks in CATATAN:
        r += 1
        ws.cell(r, 1, f"- {teks}").font = f_catatan
    atur_cetak(ws, hdr)
    ref_kat = {e["uraian"]: f"RAB!{cj}{x}" for e, x in zip([e for e in _entri_rab(data["baris"]) if e["jenis"] == "subtotal"], sub_rab)}

    # ---------- Sheet Rekapitulasi (KF-13) ----------
    if opsi.rekap:
        wr = wb.create_sheet("Rekapitulasi")
        for kol, lebar in zip("ABCD", (7, 60, 24, 12)):
            wr.column_dimensions[kol].width = lebar
        r = judul_sheet(wr, "REKAPITULASI RENCANA ANGGARAN BIAYA", 4)
        hdr = r
        header(wr, r, ["No", "Uraian", "Jumlah (Rp)", "Bobot"])
        r += 1
        gaya_baris(wr, r, 4, f_bold, fill_kat)
        wr.cell(r, 1, "A")
        wr.cell(r, 2, "BIAYA LANGSUNG")
        r += 1
        awalA = r
        for i, nama in enumerate(ref_kat, 1):
            gaya_baris(wr, r, 4)
            wr.cell(r, 1, _romawi(i)).alignment = Alignment(horizontal="center")
            wr.cell(r, 2, nama.split(". ", 1)[1])
            wr.cell(r, 3, f"={ref_kat[nama]}").number_format = ANGKA
            r += 1
        rA = r
        gaya_baris(wr, r, 4, f_bold)
        wr.cell(r, 2, "Jumlah A").alignment = Alignment(horizontal="right")
        wr.cell(r, 3, f"=SUM(C{awalA}:C{r - 1})").number_format = ANGKA
        for x in range(awalA, r + 1):
            wr.cell(x, 4, f"=IF($C${rA}=0,0,C{x}/$C${rA})").number_format = PERSEN
        r += 1
        gaya_baris(wr, r, 4, f_bold, fill_kat)
        wr.cell(r, 1, "B")
        wr.cell(r, 2, "BIAYA TIDAK LANGSUNG")
        r += 1
        awalB = r
        for i, b in enumerate(ring["item_tidak_langsung"], 1):
            gaya_baris(wr, r, 4)
            wr.cell(r, 1, i).alignment = Alignment(horizontal="center")
            wr.cell(r, 2, _label_btl(b))
            if b["jenis"] == "persen":
                wr.cell(r, 4, b["nilai"] / 100).number_format = PERSEN
                wr.cell(r, 4).font = f_input  # sel biru = nilai input
                wr.cell(r, 3, f"=C${rA}*D{r}")
            else:
                wr.cell(r, 3, b["nilai"]).font = f_input
            wr.cell(r, 3).number_format = ANGKA
            r += 1
        if not ring["item_tidak_langsung"]:
            gaya_baris(wr, r, 4)
            wr.cell(r, 2, "Tidak ada (diperhitungkan terpisah oleh pengguna)").font = f_catatan
            wr.cell(r, 3, 0).number_format = ANGKA
            r += 1
        rB = r
        gaya_baris(wr, r, 4, f_bold)
        wr.cell(r, 2, "Jumlah B").alignment = Alignment(horizontal="right")
        wr.cell(r, 3, f"=SUM(C{awalB}:C{r - 1})").number_format = ANGKA
        r += 2
        for teks, rumus in (
            ("JUMLAH (A + B)", f"=C{rA}+C{rB}"),
            (f"PPN {PPN_RATE:.0%}", f"=C{r}*{PPN_RATE!r}"),
            ("TOTAL", f"=C{r}+C{r + 1}"),
            ("TOTAL DIBULATKAN", f"=ROUNDDOWN(C{r + 2},-3)"),
        ):
            gaya_baris(wr, r, 4, f_bold)
            wr.cell(r, 2, teks).alignment = Alignment(horizontal="right")
            wr.cell(r, 3, rumus).number_format = ANGKA
            r += 1
        r += 1
        wr.cell(r, 1, f"Terbilang: {ring['terbilang']}").font = Font(name="Arial", size=10, italic=True, bold=True)
        r += 1
        wr.cell(r, 1, f"Biaya umum & keuntungan {BUK_RATE:.0%} sudah termasuk dalam harga satuan "
                      f"(± Rp {_rp(ring['buk'], 0)} dari biaya langsung).").font = f_catatan
        atur_cetak(wr, hdr)

    # ---------- Sheet RAB Rinci ----------
    if opsi.rinci:
        wi = wb.create_sheet("RAB Rinci")
        for i, k in enumerate(kolom, 1):
            wi.column_dimensions[L(i)].width = k[2]
        r = judul_sheet(wi, "RAB RINCI PER TIPE ELEMEN", len(kolom))
        wi.cell(r - 1, 1, KET_RAB_RINCI).font = f_catatan
        hdr = r
        header(wi, r, [k[1] for k in kolom])
        r, sub, items = tulis_tabel(wi, r + 1, kolom, _entri_rinci(data["baris"]))
        r += 1
        gaya_baris(wi, r, len(kolom), f_bold)
        wi.cell(r, cu, "JUMLAH BIAYA LANGSUNG").alignment = Alignment(horizontal="right")
        wi.cell(r, ci["jumlah"], "=" + "+".join(f"{cj}{x}" for x in sub)).number_format = ANGKA
        isi_bobot(wi, kolom, items + sub, f"${cj}${r}")
        r += 2
        wi.cell(r, 1, "Selisih terhadap sheet RAB (harus 0):").font = f_catatan
        wi.cell(r, ci["jumlah"], f"={cj}{r - 2}-RAB!{cj}{baris_A}")
        wi.cell(r, ci["jumlah"]).number_format = ANGKA
        wi.cell(r, ci["jumlah"]).font = f_catatan
        atur_cetak(wi, hdr)

    # ---------- Sheet Kebutuhan Besi ----------
    if opsi.besi and data["besi"]:
        wb_ = wb.create_sheet("Kebutuhan Besi")
        for kol, lebar in zip("ABCDEFG", (6, 12, 16, 14, 16, 16, 14)):
            wb_.column_dimensions[kol].width = lebar
        r = judul_sheet(wb_, "KEBUTUHAN BESI TULANGAN PER DIAMETER", 7)
        hdr = r
        header(wb_, r, ["No", "Diameter", "Jenis Baja", "Berat (kg/m')", "Panjang Total (m)", "Berat Total (kg)", "Batang 12 m"])
        r += 1
        awal = r
        for i, k in enumerate(data["besi"], 1):
            gaya_baris(wb_, r, 7)
            wb_.cell(r, 1, i).alignment = Alignment(horizontal="center")
            wb_.cell(r, 2, k["label"])
            wb_.cell(r, 3, k["jenis"])
            wb_.cell(r, 4, f"=PI()*{k['diameter']!r}^2/4*7850/1000000").number_format = ANGKA3
            wb_.cell(r, 6, k["berat"]).number_format = ANGKA
            wb_.cell(r, 5, f"=F{r}/D{r}").number_format = ANGKA
            wb_.cell(r, 7, f"=ROUNDUP(E{r}/12,0)")
            r += 1
        gaya_baris(wb_, r, 7, f_bold)
        wb_.cell(r, 3, "JUMLAH")
        wb_.cell(r, 6, f"=SUM(F{awal}:F{r - 1})").number_format = ANGKA
        wb_.cell(r, 7, f"=SUM(G{awal}:G{r - 1})")
        r += 2
        wb_.cell(r, 1, "Berat termasuk sisa potongan 5% (Rumus 2.33). Rincian tipe penulangan:").font = f_catatan
        for t in data["tipe"]:
            r += 1
            wb_.cell(r, 2, f"{t['label']} — {t['jumlah_elemen']} buah: {t['ringkas']}").font = f_norm
        atur_cetak(wb_, hdr)

    # ---------- Sheet Rekap per Lantai (KF-12) ----------
    rincian = _rincian_lantai(data) if (opsi.lantai or opsi.per_lantai) else []
    if opsi.lantai:
        wl = wb.create_sheet("Rekap per Lantai")
        for kol, lebar in zip("ABCDEFGH", (7, 46, 11, 14, 15, 14, 22, 10)):
            wl.column_dimensions[kol].width = lebar
        r = judul_sheet(wl, "REKAPITULASI BIAYA & KEBUTUHAN STRUKTUR PER LANTAI", 8)
        hdr = r
        header(wl, r, ["No", "Lantai / Kategori Pekerjaan", "Elemen", "Beton (m³)", "Bekisting (m²)", "Besi (kg)",
                       "Jumlah (Rp)", "Bobot"])
        r += 1
        baris_lantai = []
        for i, x in enumerate(rincian, 1):
            rl = r
            gaya_baris(wl, rl, 8, f_bold, fill_kat)
            wl.cell(rl, 1, i).alignment = Alignment(horizontal="center")
            elev = "" if x["elevasi"] is None else f" (elevasi {_rp(x['elevasi'])} m)"
            wl.cell(rl, 2, x["lantai"] + elev)
            wl.cell(rl, 3, x["jumlah_elemen"])
            r += 1
            per_kat = _struktur_per_kategori(x["baris"])
            for k in x["kategori"]:
                gaya_baris(wl, r, 8)
                wl.cell(r, 2, "    " + k["kategori"])
                for c, kunci in ((4, "beton"), (5, "bekisting"), (6, "besi")):
                    v = per_kat.get(k["kategori"], {}).get(kunci, 0.0)
                    wl.cell(r, c, v if v else None).number_format = ANGKA
                wl.cell(r, 7, k["total"]).number_format = ANGKA
                r += 1
            for c in "DEFG":
                wl[f"{c}{rl}"] = f"=SUM({c}{rl + 1}:{c}{r - 1})"
                wl[f"{c}{rl}"].number_format = ANGKA
            baris_lantai.append(rl)
        gaya_baris(wl, r, 8, f_bold)
        wl.cell(r, 2, f"JUMLAH {len(rincian)} LANTAI").alignment = Alignment(horizontal="right")
        for c in "DEFG":
            wl[f"{c}{r}"] = "=" + ("+".join(f"{c}{x}" for x in baris_lantai) or "0")
            wl[f"{c}{r}"].number_format = ANGKA
        for x in range(hdr + 1, r + 1):
            if wl.cell(x, 7).value is not None:
                wl.cell(x, 8, f"=IF($G${r}=0,0,G{x}/$G${r})").number_format = PERSEN
        wl.cell(r + 2, 2, "Selisih terhadap sheet RAB (harus 0)").font = f_catatan
        wl.cell(r + 2, 7, f"=G{r}-RAB!{cj}{baris_A}").number_format = ANGKA
        wl.cell(r + 2, 7).font = f_catatan
        atur_cetak(wl, hdr, "landscape")

        # ---------- Sheet Kebutuhan per Lantai (tabel silang, kolom = lantai) ----------
        wk = wb.create_sheet("Kebutuhan per Lantai")
        n_l = len(rincian)
        n_kol = 3 + n_l + 1
        for i, lebar in enumerate([6, 44, 8] + [14] * n_l + [16], 1):
            wk.column_dimensions[L(i)].width = lebar
        r = judul_sheet(wk, "KEBUTUHAN BESI, BETON, BAHAN, TENAGA KERJA & ALAT PER LANTAI", n_kol)
        judul_lantai = [x["lantai"] for x in rincian]
        hdr_pertama = r

        def matriks(r, judul, baris_m, fmt=ANGKA, catatan=None):
            gaya_baris(wk, r, n_kol, f_bold, fill_kat)
            wk.cell(r, 1, judul)
            wk.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n_kol)
            r += 1
            header(wk, r, ["No", "Uraian", "Sat"] + judul_lantai + ["Total"])
            r += 1
            for n, m in enumerate(baris_m, 1):
                gaya_baris(wk, r, n_kol)
                wk.cell(r, 1, n).alignment = Alignment(horizontal="center")
                wk.cell(r, 2, m["label"])
                wk.cell(r, 3, m["satuan"]).alignment = Alignment(horizontal="center")
                for j, v in enumerate(m["per_lantai"]):
                    wk.cell(r, 4 + j, v if abs(v) > 1e-12 else None).number_format = fmt
                wk.cell(r, n_kol, f"=SUM({L(4)}{r}:{L(3 + n_l)}{r})").number_format = fmt
                wk.cell(r, n_kol).font = f_bold
                r += 1
            if catatan:
                wk.cell(r, 2, catatan).font = f_catatan
                r += 1
            return r + 1

        r = matriks(r, "1. BESI TULANGAN PER DIAMETER (kg, Rumus 2.33 termasuk sisa potongan 5%)",
                    matriks_lantai(rincian, "besi"))
        mutu = {}
        for j, x in enumerate(rincian):
            for m in x["beton_mutu"]:
                mutu.setdefault(m["mutu"], [0.0] * n_l)[j] += m["volume"]
        beton = [{"label": f"Beton {k}", "satuan": "m³", "per_lantai": v} for k, v in mutu.items()]
        beton.append({"label": "Bekisting", "satuan": "m²", "per_lantai": [ringkas_lantai(x)["bekisting"] for x in rincian]})
        r = matriks(r, "2. BETON PER MUTU & BEKISTING", beton)
        for no, (tipe, nama) in enumerate(TIPE_SUMBER_DAYA, 3):
            r = matriks(r, f"{no}. {nama.upper()} (volume pekerjaan × koefisien AHSP)", matriks_lantai(rincian, tipe),
                        catatan="Kuantitas mengikuti koefisien analisa harga satuan HSPK yang dipakai di RAB.")
        atur_cetak(wk, hdr_pertama + 1, "landscape")
        wk.print_title_rows = None  # beberapa tabel: judul kolom tiap tabel ditulis sendiri

        # ---------- Sheet Rincian Perhitungan Lantai: susunan RAP + BackUp Volume per lantai ----------
        wp = wb.create_sheet("Rincian Perhitungan Lantai")
        judul_p = ["No", "Uraian Pekerjaan", "Panjang (m)", "Lebar (m)", "Tinggi / Tebal (m)", "Luas (m²)", "Jlh Unit",
                   "Volume / Unit", "Total Volume", "Sat", "Harga Satuan (Rp)", "Jumlah Harga (Rp)", "Kode / Keterangan"]
        for i, lebar in enumerate((6, 52, 11, 10, 11, 11, 9, 13, 14, 7, 16, 18, 44), 1):
            wp.column_dimensions[L(i)].width = lebar
        r = judul_sheet(wp, "RINCIAN PERHITUNGAN PER LANTAI: RAB & BACKUP VOLUME (DARI LANTAI TERBAWAH SAMPAI PENUTUP)", 13)
        hdr_p = r
        header(wp, r, judul_p)
        r += 1
        f_backup = Font(name="Arial", size=9, italic=True, color="555555")
        bagian_rows, item_rows, lantai_row, total_lantai = [], [], None, None

        def tutup_bagian():
            if bagian_rows and item_rows:
                rb = bagian_rows[-1]
                wp[f"L{rb}"] = "=" + "+".join(f"L{x}" for x in item_rows)
                wp[f"L{rb}"].number_format = ANGKA

        def tutup_lantai():
            nonlocal r
            tutup_bagian()
            if lantai_row is None:
                return
            gaya_baris(wp, r, 13, f_bold)
            wp.cell(r, 2, "JUMLAH BIAYA LANGSUNG LANTAI INI").alignment = Alignment(horizontal="right")
            wp[f"L{r}"] = "=" + ("+".join(f"L{x}" for x in semua_bagian) or "0")
            wp[f"L{r}"].number_format = ANGKA
            r += 1
            wp.cell(r, 2, "Selisih terhadap rekap per lantai (harus 0)").font = f_catatan
            wp[f"L{r}"] = f"=L{r - 1}-{total_lantai!r}"
            wp[f"L{r}"].number_format = ANGKA
            wp[f"L{r}"].font = f_catatan
            r += 1

        semua_bagian = []
        for e in entri_rincian_perhitungan(rincian):
            jenis = e[0]
            if jenis == "lantai":
                if lantai_row is not None:
                    tutup_lantai()  # jumlah lantai sebelumnya
                    r += 1
                lantai_row, total_lantai, semua_bagian, bagian_rows, item_rows = r, e[3], [], [], []
                gaya_baris(wp, r, 13, Font(name="Arial", size=11, bold=True), fill_head)
                wp.cell(r, 2, e[1])
                wp.cell(r, 13, e[2])
                r += 1
            elif jenis == "bagian":
                tutup_bagian()
                item_rows = []
                gaya_baris(wp, r, 13, f_bold, fill_kat)
                wp.cell(r, 1, e[1])
                wp.cell(r, 2, e[2])
                bagian_rows.append(r)
                semua_bagian.append(r)
                r += 1
            elif jenis == "grup":
                gaya_baris(wp, r, 13, f_bold, fill_grup)
                wp.cell(r, 1, e[1]).alignment = Alignment(horizontal="center")
                wp.cell(r, 2, e[2])
                r += 1
            elif jenis == "item":
                _, no, uraian, volume, satuan, harga, jumlah, kode, backup = e
                ri = r
                gaya_baris(wp, ri, 13)
                wp.cell(ri, 1, no).alignment = Alignment(horizontal="center", vertical="top")
                wp.cell(ri, 2, ("    – " if no == "-" else "") + uraian).alignment = Alignment(wrap_text=True, vertical="top")
                wp.cell(ri, 10, satuan).alignment = Alignment(horizontal="center")
                wp.cell(ri, 11, harga).number_format = ANGKA
                wp.cell(ri, 12, f"=I{ri}*K{ri}").number_format = ANGKA
                wp.cell(ri, 13, kode)
                r += 1
                awal = r
                for b_ in backup:
                    gaya_baris(wp, r, 13, f_backup)
                    wp.cell(r, 2, f"      • {b_['uraian']}").alignment = Alignment(wrap_text=True, vertical="top")
                    for c, v in ((3, b_["panjang"]), (4, b_["lebar"]), (5, b_["tinggi"]), (6, b_["luas"])):
                        if v:
                            wp.cell(r, c, v).number_format = ANGKA
                    wp.cell(r, 7, b_["unit"]).alignment = Alignment(horizontal="center")
                    wp.cell(r, 8, b_["per_unit"]).number_format = "#,##0.0000"
                    wp.cell(r, 9, f"=G{r}*H{r}").number_format = "#,##0.000"
                    wp.cell(r, 10, satuan).alignment = Alignment(horizontal="center")
                    r += 1
                wp.cell(ri, 9, f"=SUM(I{awal}:I{r - 1})" if r > awal else volume).number_format = "#,##0.000"
                wp.cell(ri, 9).font = f_bold
                item_rows.append(ri)
            elif jenis == "sub":
                tutup_bagian()
                bagian_rows, item_rows = [], []
                gaya_baris(wp, r, 13, f_bold, fill_grup)
                wp.cell(r, 2, e[1])
                r += 1
            elif jenis == "kebutuhan":
                _, no, uraian, volume, satuan, harga, jumlah, ket = e
                gaya_baris(wp, r, 13)
                wp.cell(r, 1, no).alignment = Alignment(horizontal="center")
                wp.cell(r, 2, uraian)
                wp.cell(r, 9, volume).number_format = ANGKA
                wp.cell(r, 10, satuan).alignment = Alignment(horizontal="center")
                if harga is not None:
                    wp.cell(r, 11, harga).number_format = ANGKA
                    wp.cell(r, 12, f"=I{r}*K{r}").number_format = ANGKA
                wp.cell(r, 13, ket).alignment = Alignment(wrap_text=True, vertical="top")
                r += 1
            else:  # penutup bangunan: setelah jumlah lantai teratas
                tutup_lantai()
                lantai_row = None
                r += 1
                gaya_baris(wp, r, 13, f_bold, fill_kat)
                wp.cell(r, 2, "PENUTUP BANGUNAN" + (f" ({e[1]})" if e[1] else ""))
                if not e[2]:
                    r += 1
                    gaya_baris(wp, r, 13)
                    wp.cell(r, 2, "Model tidak memuat atap, dak beton, maupun plafon.")
                for jenis_p, uraian in e[2]:
                    r += 1
                    gaya_baris(wp, r, 13)
                    wp.cell(r, 2, jenis_p).font = f_bold
                    wp.cell(r, 13, uraian).alignment = Alignment(wrap_text=True, vertical="top")
                r += 1
        tutup_lantai()
        atur_cetak(wp, hdr_p, "landscape")
        wk.freeze_panes = "D1"  # kolom No, Uraian, Sat tetap terlihat saat digeser ke kanan (banyak lantai)

    # ---------- Sheet per lantai (satu sheet tiap lantai, berapa pun jumlah lantainya) ----------
    if opsi.per_lantai:
        terpakai = {ws_.title.lower() for ws_ in wb.worksheets}
        for i, x in enumerate(rincian, 1):
            _sheet_lantai(wb.create_sheet(nama_sheet_lantai(i, x["lantai"], terpakai)), x, i, len(rincian),
                          judul_sheet, header, gaya_baris, atur_cetak, L,
                          dict(f_bold=f_bold, f_catatan=f_catatan, fill_kat=fill_kat, fill_grup=fill_grup,
                               ANGKA=ANGKA, ANGKA3=ANGKA3, PERSEN=PERSEN, Alignment=Alignment, Font=Font))

    # ---------- Sheet Detail Elemen ----------
    if opsi.detail:
        wd = wb.create_sheet("Detail Elemen")
        kd = _kolom_tabel(opsi, "detail", any(r.get("catatan") for r in data["baris"]))
        cd = {k[0]: i + 1 for i, k in enumerate(kd)}
        for i, k in enumerate(kd, 1):
            wd.column_dimensions[L(i)].width = k[2]
        r = judul_sheet(wd, "DETAIL VOLUME PER ELEMEN & LANTAI", len(kd))
        hdr = r
        header(wd, r, [k[1] for k in kd])
        r += 1
        awal = r
        for n, d in enumerate(baris_detail(data["baris"]), 1):
            gaya_baris(wd, r, len(kd))
            nilai = {**d, "no": n, "uraian": d["pekerjaan"]}
            for kunci, idx in cd.items():
                if kunci == "jumlah":
                    continue
                s = wd.cell(r, idx, nilai[kunci])
                if kunci in ("volume", "harga"):
                    s.number_format = ANGKA
                if kunci in ("uraian", "rumus", "catatan"):
                    s.alignment = Alignment(wrap_text=True, vertical="top")
            if "harga" in cd:
                wd.cell(r, cd["jumlah"], f"={L(cd['volume'])}{r}*{L(cd['harga'])}{r}")
            else:
                wd.cell(r, cd["jumlah"], d["jumlah"])
            wd.cell(r, cd["jumlah"]).number_format = ANGKA
            r += 1
        gaya_baris(wd, r, len(kd), f_bold)
        wd.cell(r, cd["uraian"], "JUMLAH").alignment = Alignment(horizontal="right")
        cdj = L(cd["jumlah"])
        wd.cell(r, cd["jumlah"], f"=SUM({cdj}{awal}:{cdj}{r - 1})").number_format = ANGKA
        r += 2
        wd.cell(r, cd["uraian"], "Selisih terhadap sheet RAB (harus 0)").font = f_catatan
        wd.cell(r, cd["jumlah"], f"={cdj}{r - 2}-RAB!{cj}{baris_A}").number_format = ANGKA
        wd.cell(r, cd["jumlah"]).font = f_catatan
        atur_cetak(wd, hdr, "landscape")

    path = tulis_aman(path, lambda tmp: wb.save(tmp))
    catat("export", f"Export Excel: {path} (isi: {_isi(opsi)}; kolom: {', '.join(opsi.kolom) or '-'})", data.get("proyek_id"))
    return str(path)


def entri_rincian_perhitungan(rincian: list) -> list:
    """Rincian perhitungan per lantai berurutan (sama dengan panel aplikasi), susunan RAP + BackUp Volume:
    ('lantai', judul, keterangan, biaya) | ('bagian', no, judul, total) | ('grup', no, judul, total)
    | ('item', no, uraian, volume, satuan, harga, jumlah, kode, backup) | ('sub', judul)
    | ('kebutuhan', no, uraian, volume, satuan, harga_atau_None, jumlah_atau_None, keterangan)
    | ('penutup', lantai, [(jenis, uraian)])."""
    out = []
    n = len(rincian)
    lantai_penutup, unsur_penutup = penutup_bangunan(rincian)
    for i, x in enumerate(rincian, 1):
        elev = "" if x["elevasi"] is None else f", elevasi {_rp(x['elevasi'])} m"
        out.append(("lantai", x["lantai"].upper(), f"Lantai {i} dari {n}{elev}, {x['jumlah_elemen']} elemen", x["total"]))
        for bag in rab_lantai(x):
            out.append(("bagian", bag["no"], bag["judul"], bag["total"]))
            for g in bag["grup"]:
                if g["judul"]:
                    out.append(("grup", g["no"], g["judul"], g["total"]))
                for it in g["items"]:
                    out.append(("item", it["no"], it["uraian"], it["volume"], it["satuan"], it["harga"], it["jumlah"],
                                it["kode"], it["backup"]))
        if x["besi"] or x["besi_rasio"] > 1e-9:
            out.append(("sub", "KEBUTUHAN BESI TULANGAN PER DIAMETER (Rumus 2.33, termasuk sisa potongan 5%)"))
            for n_, k in enumerate(x["besi"], 1):
                out.append(("kebutuhan", str(n_), f"Besi {k['label']} {k['jenis']}", k["berat"], "kg", None, None,
                            f"{_rp(k['berat'])} kg ÷ {_rp(k['berat_per_m'], 3)} kg/m' = {_rp(k['panjang'], 1)} m' "
                            f"→ {k['batang']} batang @ 12 m"))
            if x["besi_rasio"] > 1e-9:
                out.append(("kebutuhan", "", "Besi tanpa rincian diameter (asumsi rasio kg/m³)", x["besi_rasio"], "kg",
                            None, None, "elemen tanpa tipe penulangan"))
        for tipe, nama in TIPE_SUMBER_DAYA:
            items = x["sumber_daya"][tipe]
            if not items:
                continue
            out.append(("sub", f"KEBUTUHAN {nama.upper()} (VOLUME × KOEFISIEN AHSP, HARGA DASAR)"))
            for n_, d in enumerate(items, 1):
                ket = (f"{_rp(d['jumlah'])} kg ÷ {BERAT_ZAK_SEMEN} kg = {d['keterangan'].replace('≈ ', '')}"
                       if "zak" in d["keterangan"] else d["keterangan"])
                out.append(("kebutuhan", str(n_), d["nama"], d["jumlah"], d["satuan"], d["harga"], d["biaya"], ket))
        if x["lantai"] == lantai_penutup:
            out.append(("penutup", x["lantai"], unsur_penutup))
    if lantai_penutup is None:
        out.append(("penutup", None, []))
    return out


def _struktur_per_kategori(baris: list) -> dict:
    """{kategori: {'beton', 'bekisting', 'besi'}} untuk kolom rekap per lantai."""
    hasil = {}
    for r in baris:
        kode = r.get("kode_ahsp") or ""
        kunci = ("beton" if kode.startswith(("BTN.", "LTK.")) else "bekisting" if kode.startswith("BSK.")
                 else "besi" if kode.startswith("BSI.") else None)
        if kunci:
            d = hasil.setdefault(r["kategori"], {"beton": 0.0, "bekisting": 0.0, "besi": 0.0})
            d[kunci] += r["volume_pekerjaan"] or 0.0
    return hasil


def _sheet_lantai(ws, x, i, n_lantai, judul_sheet, header, gaya_baris, atur_cetak, L, g):
    """Satu sheet rincian lantai. Kolom: A No | B Uraian | C-G angka (judul kolom per bagian) | H Jumlah."""
    Alignment, f_bold, f_catatan = g["Alignment"], g["f_bold"], g["f_catatan"]
    ANGKA, PERSEN = g["ANGKA"], g["PERSEN"]
    for kol, lebar in zip("ABCDEFGH", (6, 48, 14, 13, 15, 16, 14, 20)):
        ws.column_dimensions[kol].width = lebar
    r = judul_sheet(ws, f"RINCIAN KEBUTUHAN LANTAI {i} DARI {n_lantai}: {x['lantai'].upper()}", 8)
    elev = "-" if x["elevasi"] is None else f"{_rp(x['elevasi'])} m"
    ws.cell(r - 1, 1, f"Lantai : {x['lantai']}   |   Elevasi : {elev}   |   Jumlah elemen : {x['jumlah_elemen']}").font = f_bold
    r += 1
    hdr_awal = r

    def bagian(r, judul, kolom):
        gaya_baris(ws, r, 8, f_bold, g["fill_kat"])
        ws.cell(r, 1, judul.split(".", 1)[0])
        ws.cell(r, 2, judul.split(". ", 1)[1])
        header(ws, r + 1, kolom)
        return r + 2

    def penutup(r, awal, teks, kolom_jumlah="H"):
        gaya_baris(ws, r, 8, f_bold)
        ws.cell(r, 2, teks).alignment = Alignment(horizontal="right")
        ws[f"{kolom_jumlah}{r}"] = f"=SUM({kolom_jumlah}{awal}:{kolom_jumlah}{r - 1})" if r > awal else 0
        ws[f"{kolom_jumlah}{r}"].number_format = ANGKA
        return r

    # I. biaya per kategori
    r = bagian(r, "I. REKAP BIAYA PER KATEGORI PEKERJAAN", ["No", "Kategori Pekerjaan", "", "", "", "", "Bobot", "Jumlah (Rp)"])
    awal = r
    for n, k in enumerate(x["kategori"], 1):
        gaya_baris(ws, r, 8)
        ws.cell(r, 1, n).alignment = Alignment(horizontal="center")
        ws.cell(r, 2, k["kategori"])
        ws.cell(r, 8, k["total"]).number_format = ANGKA
        r += 1
    r_total = penutup(r, awal, "JUMLAH BIAYA LANTAI INI")
    for x_ in range(awal, r_total + 1):
        ws.cell(x_, 7, f"=IF($H${r_total}=0,0,H{x_}/$H${r_total})").number_format = PERSEN
    r += 2

    # II. struktur per tipe
    if x["struktur"]:
        r = bagian(r, "II. STRUKTUR BETON PER TIPE ELEMEN",
                   ["No", "Tipe Elemen (mutu beton)", "Jumlah (bh)", "Beton (m³)", "Bekisting (m²)", "Besi (kg)",
                    "Rasio (kg/m³)", "Biaya (Rp)"])
        awal = r
        for n, s_ in enumerate(x["struktur"], 1):
            gaya_baris(ws, r, 8)
            ws.cell(r, 1, n).alignment = Alignment(horizontal="center")
            ws.cell(r, 2, f"{s_['label']} ({s_['mutu']})" + ("" if rasio_wajar(s_["rasio"]) else " — periksa rasio besi"))
            ws.cell(r, 3, s_["jumlah_elemen"] or None)
            for c, v in ((4, s_["beton"]), (5, s_["bekisting"]), (6, s_["besi"])):
                ws.cell(r, c, v if v else None).number_format = ANGKA
            ws.cell(r, 7, f'=IF(N(D{r})=0,"-",F{r}/D{r})').number_format = "#,##0"
            ws.cell(r, 8, s_["biaya"]).number_format = ANGKA
            r += 1
        gaya_baris(ws, r, 8, f_bold)
        ws.cell(r, 2, "JUMLAH STRUKTUR").alignment = Alignment(horizontal="right")
        for c in "CDEFH":
            ws[f"{c}{r}"] = f"=SUM({c}{awal}:{c}{r - 1})"
            ws[f"{c}{r}"].number_format = ANGKA if c != "C" else "0"
        ws[f"G{r}"] = f'=IF(N(D{r})=0,"-",F{r}/D{r})'
        ws[f"G{r}"].number_format = "#,##0"
        r += 1
        for m in x["beton_mutu"]:
            ws.cell(r, 2, f"Total beton {m['mutu']}").font = f_catatan
            ws.cell(r, 4, m["volume"]).number_format = ANGKA
            r += 1
        r += 1

    # III. besi per diameter
    if x["besi"] or x["besi_rasio"] > 1e-9:
        r = bagian(r, "III. KEBUTUHAN BESI TULANGAN PER DIAMETER",
                   ["No", "Diameter / Jenis Baja", "Berat (kg/m')", "Panjang (m')", "Batang 12 m", "Berat (kg)", "", ""])
        awal = r
        for n, k in enumerate(x["besi"], 1):
            gaya_baris(ws, r, 8)
            ws.cell(r, 1, n).alignment = Alignment(horizontal="center")
            ws.cell(r, 2, f"{k['label']}  {k['jenis']}")
            ws.cell(r, 3, f"=PI()*{k['diameter']!r}^2/4*7850/1000000").number_format = g["ANGKA3"]
            ws.cell(r, 6, k["berat"]).number_format = ANGKA
            ws.cell(r, 4, f"=F{r}/C{r}").number_format = ANGKA
            ws.cell(r, 5, f"=ROUNDUP(D{r}/12,0)")
            r += 1
        if x["besi_rasio"] > 1e-9:
            gaya_baris(ws, r, 8)
            ws.cell(r, 2, "Besi tanpa rincian diameter (asumsi rasio kg/m³)")
            ws.cell(r, 6, x["besi_rasio"]).number_format = ANGKA
            r += 1
        gaya_baris(ws, r, 8, f_bold)
        ws.cell(r, 2, "JUMLAH BESI").alignment = Alignment(horizontal="right")
        ws[f"E{r}"] = f"=SUM(E{awal}:E{r - 1})"
        ws[f"F{r}"] = f"=SUM(F{awal}:F{r - 1})"
        ws[f"F{r}"].number_format = ANGKA
        r += 2

    # IV - VI. bahan, tenaga kerja, alat
    romawi = iter(("IV", "V", "VI"))
    for tipe, nama in TIPE_SUMBER_DAYA:
        items = x["sumber_daya"][tipe]
        if not items:
            continue
        r = bagian(r, f"{next(romawi)}. KEBUTUHAN {nama.upper()} (KOEFISIEN AHSP)",
                   ["No", "Uraian", "Volume", "Sat", "Keterangan", "Harga Dasar", "", "Jumlah (Rp)"])
        awal = r
        for n, d in enumerate(items, 1):
            gaya_baris(ws, r, 8)
            ws.cell(r, 1, n).alignment = Alignment(horizontal="center")
            ws.cell(r, 2, d["nama"])
            ws.cell(r, 3, d["jumlah"]).number_format = ANGKA
            ws.cell(r, 4, d["satuan"]).alignment = Alignment(horizontal="center")
            ws.cell(r, 5, d["keterangan"])
            ws.cell(r, 6, d["harga"]).number_format = ANGKA
            ws.cell(r, 8, f"=C{r}*F{r}").number_format = ANGKA
            r += 1
        penutup(r, awal, f"JUMLAH {nama.upper()}")
        r += 2
    if x["sumber_daya"]["tanpa_analisa"]:
        ws.cell(r, 2, "Pekerjaan tanpa analisa (tidak masuk kebutuhan): " + ", ".join(x["sumber_daya"]["tanpa_analisa"])).font = f_catatan
        r += 1
    ws.cell(r, 2, "Harga dasar belum termasuk biaya umum & keuntungan; kuantitas = volume pekerjaan × koefisien "
                  "analisa HSPK.").font = f_catatan
    r += 2

    # VII. RAB rinci lantai
    r = bagian(r, "VII. RINCIAN PEKERJAAN (RAB RINCI LANTAI)",
               ["No", "Uraian Pekerjaan", "Volume", "Sat", "Kode Analisa", "Harga Satuan", "", "Jumlah (Rp)"])
    sub = []
    for kn, k in enumerate(x["rinci"], 1):
        gaya_baris(ws, r, 8, f_bold, g["fill_kat"])
        ws.cell(r, 1, _romawi(kn))
        ws.cell(r, 2, k["kategori"].upper())
        r += 1
        awal = r
        nomor = 0
        for gr in k["grup"]:
            if gr["judul"]:
                nomor += 1
                gaya_baris(ws, r, 8, f_bold, g["fill_grup"])
                ws.cell(r, 1, nomor).alignment = Alignment(horizontal="center")
                ws.cell(r, 2, gr["judul"])
                r += 1
            for it in gr["items"]:
                gaya_baris(ws, r, 8)
                if not gr["judul"]:
                    nomor += 1
                    ws.cell(r, 1, nomor).alignment = Alignment(horizontal="center")
                ws.cell(r, 2, ("    – " if gr["judul"] else "") + it["label"])
                ws.cell(r, 3, it["volume"]).number_format = ANGKA
                ws.cell(r, 4, it["satuan"]).alignment = Alignment(horizontal="center")
                ws.cell(r, 5, it["kode"])
                ws.cell(r, 6, it["harga"]).number_format = ANGKA
                ws.cell(r, 8, f"=C{r}*F{r}").number_format = ANGKA
                r += 1
        sub.append(penutup(r, awal, f"Jumlah {k['kategori']}"))
        r += 1
    gaya_baris(ws, r, 8, f_bold)
    ws.cell(r, 2, "JUMLAH RAB LANTAI INI").alignment = Alignment(horizontal="right")
    ws[f"H{r}"] = "=" + ("+".join(f"H{x_}" for x_ in sub) or "0")
    ws[f"H{r}"].number_format = ANGKA
    ws.cell(r + 1, 2, "Selisih terhadap rekap kategori di atas (harus 0)").font = f_catatan
    ws[f"H{r + 1}"] = f"=H{r}-H{r_total}"
    ws[f"H{r + 1}"].number_format = ANGKA
    ws[f"H{r + 1}"].font = f_catatan
    atur_cetak(ws, hdr_awal + 1)
    ws.print_title_rows = None
    ws.freeze_panes = None


# ---------------------------------------------------------------- PDF


def export_pdf(path, data: dict, meta: dict, opsi: OpsiExport | None = None, catat_log: bool = True) -> str:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    opsi = opsi or OpsiExport()
    ukuran = landscape(A4) if opsi.orientasi_pdf == "landscape" else A4
    lebar_isi = ukuran[0] - 2.8 * cm
    nama = meta.get("nama_proyek") or data["nama_proyek"]
    ring = data["ringkasan"]
    total_A = ring["langsung"] or 1.0

    s_judul = ParagraphStyle("judul", fontName="Helvetica-Bold", fontSize=13, alignment=1, spaceAfter=6)
    s_info = ParagraphStyle("info", fontName="Helvetica", fontSize=9, leading=12)
    s_sel = ParagraphStyle("sel", fontName="Helvetica", fontSize=8, leading=10)
    s_selb = ParagraphStyle("selb", parent=s_sel, fontName="Helvetica-Bold")
    s_head = ParagraphStyle("head", parent=s_selb, alignment=1)
    s_kanan = ParagraphStyle("kanan", parent=s_selb, alignment=2)
    s_cat = ParagraphStyle("cat", fontName="Helvetica-Oblique", fontSize=7.5, leading=9, textColor=colors.HexColor("#444444"))
    s_mono = ParagraphStyle("mono", fontName="Courier", fontSize=6.5, leading=8)

    abu, biru, muda = colors.HexColor("#D9D9D9"), colors.HexColor("#EEF3F8"), colors.HexColor("#F7F9FB")
    dasar = [
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#808080")),
        ("BACKGROUND", (0, 0), (-1, 0), abu),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ("FONTSIZE", (0, 1), (-1, -1), 8),
    ]

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(1.4 * cm, 0.9 * cm, f"{nama}  |  Dibuat dengan CostStruct, {datetime.now():%d-%m-%Y %H:%M}")
        canvas.drawRightString(ukuran[0] - 1.4 * cm, 0.9 * cm, f"Halaman {doc.page}")
        canvas.restoreState()

    def kop(teks):
        out = [Paragraph(teks, s_judul)]
        for label, nilai in _info_dokumen(data, meta):
            out.append(Paragraph(f"<b>{label}</b> : {_esc(nilai)}", s_info))
        out.append(Spacer(1, 8))
        return out

    def lebar(kolom):
        total = sum(k[3] for k in kolom)
        return [lebar_isi * k[3] / total for k in kolom]

    def tabel_rab(entri, kolom):
        ci = {k[0]: i for i, k in enumerate(kolom)}
        isi = [[Paragraph(k[1], s_head) for k in kolom]]
        gaya = list(dasar)
        n = len(kolom)
        for e in entri:
            baris = [""] * n
            i = len(isi)
            if e["jenis"] in ("kategori", "grup"):
                st = s_selb
                if "no" in ci:
                    baris[ci["no"]] = Paragraph(e["no"], st)
                    baris[ci["uraian"]] = Paragraph(_esc(e["uraian"]), st)
                else:
                    baris[ci["uraian"]] = Paragraph(_esc(f"{e['no']} {e['uraian']}"), st)
                gaya.append(("BACKGROUND", (0, i), (-1, i), biru if e["jenis"] == "kategori" else muda))
            elif e["jenis"] == "item":
                if "no" in ci:
                    baris[ci["no"]] = e["no"]
                if "kode" in ci:
                    baris[ci["kode"]] = Paragraph(_esc(e["kode"]), s_sel)
                pref = "&nbsp;&nbsp;&nbsp;&nbsp;" if e.get("indent") else ""
                baris[ci["uraian"]] = Paragraph(pref + _esc(e["uraian"]), s_sel)
                baris[ci["volume"]] = _rp(e["volume"])
                baris[ci["satuan"]] = e["satuan"]
                if "harga" in ci:
                    baris[ci["harga"]] = _rp(e["harga"])
                baris[ci["jumlah"]] = _rp(e["jumlah"])
                if "bobot" in ci:
                    baris[ci["bobot"]] = _persen(e["jumlah"] / total_A)
            else:  # subtotal
                baris[ci["uraian"]] = Paragraph(_esc(e["uraian"]), s_kanan)
                baris[ci["jumlah"]] = _rp(e["jumlah"])
                if "bobot" in ci:
                    baris[ci["bobot"]] = _persen(e["jumlah"] / total_A)
                gaya += [("FONTNAME", (0, i), (-1, i), "Helvetica-Bold"),
                         ("LINEABOVE", (0, i), (-1, i), 0.8, colors.black)]
            isi.append(baris)
        for kunci in ("volume", "harga", "jumlah", "bobot"):
            if kunci in ci:
                gaya.append(("ALIGN", (ci[kunci], 1), (ci[kunci], -1), "RIGHT"))
        for kunci in ("no", "satuan"):
            if kunci in ci:
                gaya.append(("ALIGN", (ci[kunci], 1), (ci[kunci], -1), "CENTER"))
        t = Table(isi, colWidths=lebar(kolom), repeatRows=1)
        t.setStyle(TableStyle(gaya))
        return t

    def tabel_ringkas(baris, tebal_dari=0):
        t = Table(baris, colWidths=[lebar_isi - 4.2 * cm, 4.2 * cm], hAlign="RIGHT")
        t.setStyle(TableStyle([
            ("FONTNAME", (0, tebal_dari), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ALIGN", (0, 0), (-1, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#808080")),
            ("BACKGROUND", (0, -1), (-1, -1), abu),
        ]))
        return t

    kolom = _kolom_tabel(opsi)
    story = kop("RENCANA ANGGARAN BIAYA (RAB)") + [Paragraph(_esc(KET_RAB), s_cat), Spacer(1, 4)]
    story.append(tabel_rab(_entri_rab(data["baris"]), kolom))
    story.append(Spacer(1, 8))
    story.append(tabel_ringkas([
        ["A. JUMLAH BIAYA LANGSUNG", _rp(ring["langsung"])],
        ["B. BIAYA TIDAK LANGSUNG", _rp(ring["tidak_langsung"])],
        ["JUMLAH (A + B)", _rp(ring["jumlah"])],
        [f"PPN {PPN_RATE:.0%}", _rp(ring["ppn"])],
        ["TOTAL", _rp(ring["total"])],
        ["TOTAL DIBULATKAN", _rp(ring["dibulatkan"], 0)],
    ]))
    story.append(Spacer(1, 6))
    story.append(Paragraph(f"<b><i>Terbilang: {_esc(ring['terbilang'])}</i></b>", s_info))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Catatan:", s_cat))
    for teks in CATATAN:
        story.append(Paragraph(f"- {_esc(teks)}", s_cat))

    if opsi.rekap:
        story += [PageBreak()] + kop("REKAPITULASI RENCANA ANGGARAN BIAYA")
        rb = [[Paragraph(t, s_head) for t in ("No", "Uraian", "Jumlah (Rp)", "Bobot")]]
        gaya = list(dasar) + [("ALIGN", (2, 1), (3, -1), "RIGHT"), ("ALIGN", (0, 1), (0, -1), "CENTER")]

        def tambah(baris, tebal=False, latar=None):
            rb.append(baris)
            i = len(rb) - 1
            if tebal:
                gaya.append(("FONTNAME", (0, i), (-1, i), "Helvetica-Bold"))
            if latar:
                gaya.append(("BACKGROUND", (0, i), (-1, i), latar))

        tambah(["A", "BIAYA LANGSUNG", "", ""], True, biru)
        for i, k in enumerate(kelompokkan(data["baris"]), 1):
            tambah([_romawi(i), k["kategori"], _rp(k["total"]), _persen(k["total"] / total_A)])
        tambah(["", "Jumlah A", _rp(ring["langsung"]), "100,00%"], True)
        tambah(["B", "BIAYA TIDAK LANGSUNG", "", ""], True, biru)
        for i, b in enumerate(ring["item_tidak_langsung"], 1):
            tambah([str(i), Paragraph(_esc(_label_btl(b)), s_sel), _rp(b["jumlah"]), ""])
        if not ring["item_tidak_langsung"]:
            tambah(["", "Tidak ada (diperhitungkan terpisah oleh pengguna)", _rp(0), ""])
        tambah(["", "Jumlah B", _rp(ring["tidak_langsung"]), ""], True)
        for teks, nilai in (("JUMLAH (A + B)", ring["jumlah"]), (f"PPN {PPN_RATE:.0%}", ring["ppn"]),
                            ("TOTAL", ring["total"])):
            tambah(["", teks, _rp(nilai), ""], True)
        tambah(["", "TOTAL DIBULATKAN", _rp(ring["dibulatkan"], 0), ""], True, abu)
        t = Table(rb, colWidths=[1.2 * cm, lebar_isi - 8.4 * cm, 4.6 * cm, 2.6 * cm], repeatRows=1)
        t.setStyle(TableStyle(gaya))
        story += [t, Spacer(1, 6), Paragraph(f"<b><i>Terbilang: {_esc(ring['terbilang'])}</i></b>", s_info),
                  Spacer(1, 4),
                  Paragraph(f"Biaya umum & keuntungan {BUK_RATE:.0%} sudah termasuk dalam harga satuan "
                            f"(± Rp {_rp(ring['buk'], 0)} dari biaya langsung).", s_cat)]

    if opsi.rinci:
        story += [PageBreak()] + kop("RAB RINCI PER TIPE ELEMEN") + [Paragraph(_esc(KET_RAB_RINCI), s_cat), Spacer(1, 4)]
        story.append(tabel_rab(_entri_rinci(data["baris"]), kolom))

    if opsi.besi and data["besi"]:
        story += [PageBreak()] + kop("KEBUTUHAN BESI TULANGAN PER DIAMETER")
        bb = [[Paragraph(t, s_head) for t in ("No", "Diameter", "Jenis Baja", "Berat (kg/m')", "Panjang (m)", "Berat (kg)", "Batang 12 m")]]
        for i, k in enumerate(data["besi"], 1):
            bb.append([str(i), k["label"], k["jenis"], _rp(k["berat_per_m"], 3), _rp(k["panjang"]), _rp(k["berat"]), str(k["batang"])])
        bb.append(["", "", "JUMLAH", "", "", _rp(sum(k["berat"] for k in data["besi"])), str(sum(k["batang"] for k in data["besi"]))])
        t = Table(bb, colWidths=[lebar_isi * x for x in (0.07, 0.12, 0.2, 0.14, 0.16, 0.16, 0.15)], repeatRows=1)
        t.setStyle(TableStyle(dasar + [("ALIGN", (3, 1), (-1, -1), "RIGHT"), ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold")]))
        story.append(t)
        story.append(Spacer(1, 6))
        story.append(Paragraph("Berat termasuk sisa potongan 5% (Rumus 2.33). Tipe penulangan:", s_cat))
        for tp in data["tipe"]:
            story.append(Paragraph(_esc(f"• {tp['label']} — {tp['jumlah_elemen']} buah: {tp['ringkas']}"), s_info))

    rincian = _rincian_lantai(data) if (opsi.lantai or opsi.per_lantai) else []

    def tabel(isi, lebar_rel, gaya_tambahan=(), kanan_dari=None):
        """Tabel standar: baris pertama judul kolom; lebar_rel = proporsi lebar tiap kolom."""
        isi = [[Paragraph(_esc(h), s_head) for h in isi[0]]] + isi[1:]
        g = list(dasar) + list(gaya_tambahan)
        if kanan_dari is not None:
            g.append(("ALIGN", (kanan_dari, 1), (-1, -1), "RIGHT"))
        t = Table(isi, colWidths=[lebar_isi * x / sum(lebar_rel) for x in lebar_rel], repeatRows=1)
        t.setStyle(TableStyle(g))
        return t

    def angka(v, d=2):
        return _rp(v, d) if v else "-"

    if opsi.lantai:
        story += [PageBreak()] + kop("REKAPITULASI BIAYA & KEBUTUHAN STRUKTUR PER LANTAI")
        lb = [["No", "Lantai / Kategori Pekerjaan", "Elemen", "Beton (m³)", "Bekisting (m²)", "Besi (kg)", "Jumlah (Rp)", "Bobot"]]
        gaya = [("ALIGN", (0, 1), (0, -1), "CENTER")]
        for i, x in enumerate(rincian, 1):
            elev = "" if x["elevasi"] is None else f" (elevasi {_rp(x['elevasi'])} m)"
            ring_l = ringkas_lantai(x)
            lb.append([str(i), Paragraph(_esc(x["lantai"] + elev), s_selb), str(x["jumlah_elemen"]),
                       angka(ring_l["beton"]), angka(ring_l["bekisting"]), angka(ring_l["besi"]),
                       _rp(x["total"]), _persen(x["total"] / total_A)])
            gaya += [("BACKGROUND", (0, len(lb) - 1), (-1, len(lb) - 1), biru),
                     ("FONTNAME", (0, len(lb) - 1), (-1, len(lb) - 1), "Helvetica-Bold")]
            per_kat = _struktur_per_kategori(x["baris"])
            for k in x["kategori"]:
                pk = per_kat.get(k["kategori"], {})
                lb.append(["", Paragraph("&nbsp;&nbsp;&nbsp;&nbsp;" + _esc(k["kategori"]), s_sel), "",
                           angka(pk.get("beton")), angka(pk.get("bekisting")), angka(pk.get("besi")),
                           _rp(k["total"]), _persen(k["total"] / total_A)])
        tot = [ringkas_lantai(x) for x in rincian]
        lb.append(["", f"JUMLAH {len(rincian)} LANTAI", "", angka(sum(t_["beton"] for t_ in tot)),
                   angka(sum(t_["bekisting"] for t_ in tot)), angka(sum(t_["besi"] for t_ in tot)),
                   _rp(ring["langsung"]), "100,00%"])
        gaya.append(("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"))
        story.append(tabel(lb, (0.05, 0.31, 0.09, 0.1, 0.11, 0.1, 0.16, 0.08), gaya, kanan_dari=2))

        # tabel silang besi per diameter x lantai (lebar halaman cukup untuk <= 6 lantai)
        mb = matriks_lantai(rincian, "besi")
        if mb:
            story += [Spacer(1, 10), Paragraph("<b>Kebutuhan besi per diameter per lantai (kg)</b>", s_info), Spacer(1, 4)]
            for awal in range(0, len(rincian), 6):
                bagian_l = rincian[awal:awal + 6]
                isi = [["Diameter"] + [x["lantai"] for x in bagian_l] + (["Total"] if awal + 6 >= len(rincian) else [])]
                for m in mb:
                    baris_m = [m["label"]] + [angka(v) for v in m["per_lantai"][awal:awal + 6]]
                    if awal + 6 >= len(rincian):
                        baris_m.append(_rp(m["total"]))
                    isi.append(baris_m)
                story += [tabel(isi, [0.3] + [0.14] * (len(isi[0]) - 1), kanan_dari=1), Spacer(1, 6)]

        story += [PageBreak()] + kop("RINCIAN PERHITUNGAN PER LANTAI (RAB & BACKUP VOLUME)")
        story.append(Paragraph("Susunan RAP per lantai, berurutan dari lantai terbawah sampai penutup bangunan. Baris • adalah "
                               "backup volume: dimensi elemen × jumlah unit = volume.", s_cat))
        s_bk = ParagraphStyle("bk", parent=s_sel, fontName="Helvetica-Oblique", fontSize=7, leading=8.5,
                              textColor=colors.HexColor("#555555"))
        isi, gaya = None, []

        def tutup_tabel():
            if isi and len(isi) > 1:
                story.append(tabel(isi, (0.05, 0.275, 0.165, 0.08, 0.09, 0.055, 0.125, 0.16), gaya + [
                    ("ALIGN", (3, 1), (3, -1), "CENTER"), ("ALIGN", (5, 1), (5, -1), "CENTER"),
                    ("ALIGN", (4, 1), (4, -1), "RIGHT"), ("ALIGN", (6, 1), (-1, -1), "RIGHT")]))

        def baris(*sel, latar=None, tebal=False):
            isi.append(list(sel))
            i_ = len(isi) - 1
            if latar:
                gaya.append(("BACKGROUND", (0, i_), (-1, i_), latar))
            if tebal:
                gaya.append(("FONTNAME", (0, i_), (-1, i_), "Helvetica-Bold"))

        for e in entri_rincian_perhitungan(rincian):
            jenis = e[0]
            if jenis == "lantai":
                tutup_tabel()
                story += [Spacer(1, 8), Paragraph(f"<b>{_esc(e[1])}</b> &nbsp; <font size=8>{_esc(e[2])}; biaya "
                                                  f"langsung Rp {_rp(e[3], 0)}</font>", s_info), Spacer(1, 3)]
                isi, gaya = [["No", "Uraian Pekerjaan", "Dimensi / Keterangan", "Unit", "Volume", "Sat", "Harga Sat.",
                              "Jumlah (Rp)"]], []
            elif jenis == "bagian":
                baris(e[1], Paragraph(f"<b>{_esc(e[2])}</b>", s_sel), "", "", "", "", "", _rp(e[3]), latar=biru, tebal=True)
            elif jenis == "grup":
                baris(e[1], Paragraph(f"<b>{_esc(e[2])}</b>", s_sel), "", "", "", "", "", _rp(e[3]), latar=muda, tebal=True)
            elif jenis == "item":
                _, no, uraian, volume, satuan, harga, jumlah, kode, backup = e
                baris(no, Paragraph(("&nbsp;&nbsp;– " if no == "-" else "") + _esc(uraian), s_sel),
                      Paragraph(_esc(kode), s_sel), "", _rp(volume, 3), satuan, _rp(harga), _rp(jumlah))
                for b_ in backup:
                    dim = " × ".join(f"{h} {_rp(v)}" for h, v in (("P", b_["panjang"]), ("L", b_["lebar"]),
                                                                    ("T", b_["tinggi"])) if v)
                    if not dim and b_["luas"]:
                        dim = f"Luas {_rp(b_['luas'])} m²"
                    baris("", Paragraph("&nbsp;&nbsp;&nbsp;&nbsp;• " + _esc(b_["uraian"]), s_bk), Paragraph(_esc(dim), s_bk),
                          Paragraph(f"{b_['unit']} × {_rp(b_['per_unit'], 3)}", s_bk), Paragraph(_rp(b_["volume"], 3), s_bk),
                          Paragraph(_esc(satuan), s_bk), "", "")
            elif jenis == "sub":
                baris("", Paragraph(f"<b>{_esc(e[1])}</b>", s_sel), "", "", "", "", "", "", latar=muda)
            elif jenis == "kebutuhan":
                _, no, uraian, volume, satuan, harga, jumlah, ket = e
                baris(no, Paragraph(_esc(uraian), s_sel), Paragraph(_esc(ket), s_sel), "", _rp(volume), satuan,
                      "" if harga is None else _rp(harga), "" if jumlah is None else _rp(jumlah))
            else:
                tutup_tabel()
                isi = None
                story += [Spacer(1, 8), Paragraph("<b>PENUTUP BANGUNAN</b>" + (f" ({_esc(e[1])})" if e[1] else ""), s_info)]
                if not e[2]:
                    story.append(Paragraph("Model tidak memuat atap, dak beton, maupun plafon.", s_cat))
                for jenis_p, uraian in e[2]:
                    story.append(Paragraph(f"<b>{_esc(jenis_p)}</b>: {_esc(uraian)}", s_info))
        tutup_tabel()

    if opsi.per_lantai:
        for i, x in enumerate(rincian, 1):
            elev = "-" if x["elevasi"] is None else f"{_rp(x['elevasi'])} m"
            story += [PageBreak()] + kop(f"RINCIAN KEBUTUHAN LANTAI {i} DARI {len(rincian)}: {_esc(x['lantai'].upper())}")
            story.append(Paragraph(f"<b>Elevasi</b> : {elev} &nbsp;&nbsp; <b>Jumlah elemen</b> : {x['jumlah_elemen']}"
                                   f" &nbsp;&nbsp; <b>Biaya langsung lantai</b> : Rp {_rp(x['total'], 0)}", s_info))
            story.append(Spacer(1, 6))

            def sub_judul(teks):
                story.extend([Spacer(1, 6), Paragraph(f"<b>{_esc(teks)}</b>", s_info), Spacer(1, 3)])

            sub_judul("I. Rekap biaya per kategori pekerjaan")
            isi = [["No", "Kategori Pekerjaan", "Jumlah (Rp)", "Bobot"]]
            for n, k in enumerate(x["kategori"], 1):
                isi.append([str(n), k["kategori"], _rp(k["total"]), _persen(k["total"] / (x["total"] or 1))])
            isi.append(["", "JUMLAH", _rp(x["total"]), "100,00%"])
            story.append(tabel(isi, (0.07, 0.6, 0.2, 0.13), [("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold")], 2))

            if x["struktur"]:
                sub_judul("II. Struktur beton per tipe elemen")
                isi = [["No", "Tipe Elemen", "Mutu", "Jml", "Beton (m³)", "Bekisting (m²)", "Besi (kg)", "kg/m³", "Biaya (Rp)"]]
                gaya = []
                for n, g_ in enumerate(x["struktur"], 1):
                    isi.append([str(n), Paragraph(_esc(g_["label"]), s_sel), g_["mutu"], str(g_["jumlah_elemen"] or "-"),
                                angka(g_["beton"], 3), angka(g_["bekisting"]), angka(g_["besi"]),
                                "-" if g_["rasio"] is None else _rp(g_["rasio"], 0), _rp(g_["biaya"])])
                    if not rasio_wajar(g_["rasio"]):
                        gaya.append(("TEXTCOLOR", (7, len(isi) - 1), (7, len(isi) - 1), colors.HexColor("#B45309")))
                rs = ringkas_lantai(x)
                isi.append(["", "JUMLAH", "", "", angka(rs["beton"], 3), angka(rs["bekisting"]), angka(rs["besi"]),
                            "-" if not rs["beton"] else _rp(rs["besi"] / rs["beton"], 0),
                            _rp(sum(g_["biaya"] for g_ in x["struktur"]))])
                gaya.append(("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"))
                story.append(tabel(isi, (0.05, 0.27, 0.1, 0.05, 0.1, 0.11, 0.1, 0.07, 0.15), gaya, 3))
                story.append(Paragraph(_esc("Total beton: " + ", ".join(
                    f"{m['mutu']} {_rp(m['volume'], 3)} m³" for m in x["beton_mutu"])), s_cat))

            if x["besi"] or x["besi_rasio"] > 1e-9:
                sub_judul("III. Kebutuhan besi tulangan per diameter")
                isi = [["No", "Diameter", "Jenis Baja", "Berat (kg/m')", "Panjang (m')", "Batang 12 m", "Berat (kg)"]]
                for n, k in enumerate(x["besi"], 1):
                    isi.append([str(n), k["label"], k["jenis"], _rp(k["berat_per_m"], 3), _rp(k["panjang"]),
                                str(k["batang"]), _rp(k["berat"])])
                if x["besi_rasio"] > 1e-9:
                    isi.append(["", Paragraph("Tanpa rincian diameter (asumsi rasio kg/m³)", s_sel), "", "", "", "",
                                _rp(x["besi_rasio"])])
                isi.append(["", "JUMLAH", "", "", "", str(sum(k["batang"] for k in x["besi"])),
                            _rp(sum(k["berat"] for k in x["besi"]) + x["besi_rasio"])])
                story.append(tabel(isi, (0.06, 0.22, 0.16, 0.13, 0.15, 0.13, 0.15),
                                   [("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold")], 3))

            romawi = iter(("IV", "V", "VI"))
            for tipe, nama_t in TIPE_SUMBER_DAYA:
                items = x["sumber_daya"][tipe]
                if not items:
                    continue
                sub_judul(f"{next(romawi)}. Kebutuhan {nama_t.lower()} (volume × koefisien AHSP)")
                isi = [["No", "Uraian", "Volume", "Sat", "Keterangan", "Harga Dasar", "Jumlah (Rp)"]]
                for n, d in enumerate(items, 1):
                    isi.append([str(n), Paragraph(_esc(d["nama"]), s_sel), _rp(d["jumlah"]), d["satuan"],
                                Paragraph(_esc(d["keterangan"]), s_sel), _rp(d["harga"]), _rp(d["biaya"])])
                isi.append(["", "JUMLAH", "", "", "", "", _rp(sum(d["biaya"] for d in items))])
                story.append(tabel(isi, (0.05, 0.33, 0.12, 0.06, 0.14, 0.13, 0.17), [
                    ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"), ("ALIGN", (2, 1), (2, -1), "RIGHT"),
                    ("ALIGN", (5, 1), (-1, -1), "RIGHT"), ("ALIGN", (3, 1), (3, -1), "CENTER")]))
            if x["sumber_daya"]["tanpa_analisa"]:
                story.append(Paragraph(_esc("Pekerjaan tanpa analisa: " + ", ".join(x["sumber_daya"]["tanpa_analisa"])), s_cat))
            story.append(Paragraph("Harga dasar belum termasuk biaya umum &amp; keuntungan.", s_cat))

            sub_judul("VII. Rincian pekerjaan (RAB rinci lantai)")
            story.append(tabel_rab(_entri_rinci(x["baris"]), kolom))

    if opsi.detail:
        story += [PageBreak()] + kop("DETAIL VOLUME PER ELEMEN & LANTAI")
        kd = _kolom_tabel(opsi, "detail", any(r.get("catatan") for r in data["baris"]))
        cd = {k[0]: i for i, k in enumerate(kd)}
        db = [[Paragraph(k[1], s_head) for k in kd]]
        for n, d in enumerate(baris_detail(data["baris"]), 1):
            baris = [""] * len(kd)
            nilai = {**d, "no": str(n), "uraian": d["pekerjaan"]}
            for kunci, idx in cd.items():
                v = nilai[kunci]
                if kunci in ("volume",):
                    baris[idx] = _rp(v, 3)
                elif kunci in ("harga", "jumlah"):
                    baris[idx] = _rp(v)
                elif kunci in ("rumus",):
                    baris[idx] = Paragraph(_esc(v), s_mono)
                elif kunci in ("lantai", "elemen", "uraian", "kode", "catatan"):
                    baris[idx] = Paragraph(_esc(v), s_sel)
                else:
                    baris[idx] = v
            db.append(baris)
        gaya = list(dasar)
        for kunci in ("volume", "harga", "jumlah"):
            if kunci in cd:
                gaya.append(("ALIGN", (cd[kunci], 1), (cd[kunci], -1), "RIGHT"))
        t = Table(db, colWidths=lebar(kd), repeatRows=1)
        t.setStyle(TableStyle(gaya))
        story.append(t)

    def tulis(tmp):
        doc = SimpleDocTemplate(
            tmp, pagesize=ukuran, leftMargin=1.4 * cm, rightMargin=1.4 * cm, topMargin=1.4 * cm,
            bottomMargin=1.6 * cm, title=f"RAB {nama}", author="CostStruct",
        )
        doc.build(story, onFirstPage=footer, onLaterPages=footer)

    path = tulis_aman(path, tulis)
    if catat_log:  # pratinjau (KF-6) tidak dicatat sebagai export
        catat("export", f"Export PDF {opsi.orientasi_pdf}: {path} (isi: {_isi(opsi)})", data.get("proyek_id"))
    return str(path)


def _isi(opsi: OpsiExport) -> str:
    bagian = [("RAB", True), ("rekap", opsi.rekap), ("rinci", opsi.rinci), ("besi", opsi.besi),
              ("rekap per lantai", opsi.lantai), ("rincian per lantai", opsi.per_lantai), ("detail", opsi.detail)]
    return ", ".join(n for n, ada in bagian if ada)


def _esc(teks) -> str:
    return str(teks).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
