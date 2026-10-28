"""
Impor HSPK (Harga Satuan Pokok Kegiatan) dari workbook Excel menjadi modul data CostStruct.

    python tools/impor_hspk.py "HSPK_Kota_Bandung_2027.xlsx"

Menghasilkan src/database/hspk_bandung_2027.py berisi:
- HSD   : seluruh harga satuan dasar upah, material, dan peralatan (sheet "HSD (...)").
- AHSP  : seluruh analisa harga satuan pekerjaan bab 1-10 (katalog lengkap), lengkap dengan
          koefisien, harga komponen, dan harga satuan F dari dokumen (untuk verifikasi).
- SATUAN: satuan setiap analisa, dari daftar induk sheet "HSPK".
- BAB   : judul bab 1-10 dari daftar induk, untuk pengelompokan katalog.

Sheet khusus rumah sakit (kode bab 11 ke atas, RSKGM, MOT, Radiologi, ...) bukan HSPK standar dan
tidak diimpor.

Kolom setiap analisa dibaca dari baris header "No | Uraian | ... | Koefisien | Harga Satuan |
Jumlah Harga" milik item itu sendiri, karena susunan kolom antar sheet tidak sama.
Pemetaan kode pekerjaan aplikasi ke kode HSPK ada di src/database/seed_data.py.
"""

import argparse
import re
import sys
from pathlib import Path

import openpyxl

SHEET_HSD = {"HSD (Upah)": "upah", "HSD (Material)": "bahan", "HSD (Peralatan)": "alat"}
# Sheet yang dibaca lebih dulu. Urutannya menentukan kode ganda: analisa dengan kode yang sama di sheet
# berikutnya disimpan sebagai "kode@sheet", sehingga kode analisa yang sudah dipakai aplikasi tidak berubah.
SHEET_UTAMA = [
    "Persiapan",
    "Galian Tanah",
    "Timbunan Pemadatan",
    "Beton",
    "Pondasi",
    "Rangka Atap",
    "Penutup Atap",
    "Plafon",
    "Pasangan Dinding",
    "Plesteran Dan Acian",
    "Pengecatan dan Pelituran",
    "Penutup Lantai dan Dinding",
    "Pintu dan Jendela",
    "Kaca",
    "Kayu",
    "Struktur Kayu",
    "Besi dan Aluminium",
    "Sistem Air Limbah",
]
# Sheet ringkasan / harga dasar / daftar khusus rumah sakit: bukan analisa HSPK standar.
SHEET_LEWATI = {
    "Status", "Status Pekerjaan", "Kertas Kerja SSH (TKDN)", "SBU 2027", "ASB", "HSPK",
    "Daftar Harga Satuan Pekerjaan", "Upah Bahan", "HSD (Upah)", "HSD (Material)", "HSD (Peralatan)",
    "LIST ITEM RSKGM BARU", "LIST ITEM YG TIDAK ADA DI REGUL", "Sheet5", "III. B2. AR RS NSTD",
    "III.C.REKAP ME.RS", "RSKGM EL EK", "VI.A.SR.PARKIR", "RSKGM UTYL", "RSKGM MP", "RSKGM MP2",
    "RSKGM HEAT PUMP", "MOT", "RADIOLOGI", "ASB Rumputisasi",
}
BAB_MAKS = 10
KODE = re.compile(r"^\d+(\.\d+)+[a-zA-Z]?$")
TIPE_BLOK = {"TENAGA KERJA": "upah", "BAHAN": "bahan", "PERALATAN": "alat"}
TUJUAN = Path(__file__).resolve().parents[1] / "src" / "database" / "hspk_bandung_2027.py"


def _teks(v) -> str:
    return " ".join(str(v).split()) if v is not None else ""


def _angka(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return None


def baca_hsd(wb) -> list:
    hasil = []
    for nama_sheet, tipe in SHEET_HSD.items():
        ws = wb[nama_sheet]
        kol = None
        for row in ws.iter_rows(values_only=True):
            row = list(row) + [None] * 10
            t = [_teks(v) for v in row]
            if kol is None:
                if "URAIAN" in t:
                    i_uraian = t.index("URAIAN")
                    satuan = [i for i, v in enumerate(t) if v == "SATUAN"]
                    kol = (i_uraian, satuan[0], satuan[1])
                continue
            nama, sat, harga = t[kol[0]], t[kol[1]], _angka(row[kol[2]])
            if nama and sat and harga and harga > 0:
                hasil.append((tipe, nama, sat, round(harga, 4)))
    return hasil


def daftar_sheet(wb) -> list:
    lain = [n for n in wb.sheetnames if n not in SHEET_UTAMA and n.strip() not in SHEET_LEWATI]
    return SHEET_UTAMA + lain


def baca_ahsp(wb, sheets) -> list:
    items = []
    for sh in sheets:
        if sh not in wb.sheetnames:
            print(f"  (lewati, sheet tidak ada: {sh})")
            continue
        ws = wb[sh]
        cur, tipe, kol = None, None, None
        for row in ws.iter_rows(values_only=True):
            row = list(row) + [None] * 20
            t = [_teks(v) for v in row]
            idx = next((i for i, v in enumerate(t) if v), None)
            if idx is None:
                continue
            cs, ds = t[idx], t[idx + 1]
            kode = cs.rstrip(".")  # mis. "1.1.3.3." di dokumen
            if KODE.match(kode) and kode.count(".") >= 2 and ds and not ds.lower().startswith("jumlah"):
                cur = {"sheet": sh, "kode": kode, "uraian": ds, "komponen": [], "F": None, "BUK": None}
                items.append(cur)
                tipe, kol = None, None
                continue
            if cur is None:
                continue
            if kol is not None:
                cs, ds = t[kol["no"]], t[kol["no"] + 1]
            if cs == "No" and ds.lower().startswith("uraian"):
                low = [v.lower() for v in t]

                def cari(pred):
                    return next((i for i, v in enumerate(low) if i > idx and pred(v)), None)

                kol = {
                    "no": idx,
                    "sat": cari(lambda v: v.startswith("sat")),
                    "koef": cari(lambda v: v.startswith("koef")),
                    "harga": cari(lambda v: "harga" in v and "satuan" in v and "jumlah" not in v),
                    "jumlah": cari(lambda v: v.startswith("jumlah harga") or v == "jumlah"),
                }
                if None in kol.values():  # header tidak lengkap: analisa ini tidak bisa dibaca
                    kol, cur = None, None
                continue
            if kol is None:
                continue
            if cs in ("A", "B", "C") and ds.upper() in TIPE_BLOK:
                tipe = TIPE_BLOK[ds.upper()]
                continue
            if cs == "E" and "keuntungan" in ds.lower():
                cur["BUK"] = _angka(row[kol["harga"]])
                tipe = None
                continue
            if cs == "F" and "harga satuan" in ds.lower():
                cur["F"] = _angka(row[kol["jumlah"]])
                cur, tipe = None, None
                continue
            if tipe and ds and not ds.upper().startswith("JUMLAH"):
                koef = _angka(row[kol["koef"]])
                if koef and koef > 0:
                    cur["komponen"].append(
                        (tipe, ds, t[kol["sat"]], round(koef, 6), round(_angka(row[kol["harga"]]) or 0.0, 4))
                    )
    return [i for i in items if i["komponen"] and i["F"]]


_SATUAN_URAIAN = re.compile(r"\b1\s*(m2|m3|m'|m’|kg|buah|unit|titik|set|ls|btg|bh|ttk)(?![a-z0-9])", re.IGNORECASE)
# ejaan satuan di dokumen yang diseragamkan
_SATUAN_SAMA = {"titk": "titik", "m;'": "m'", "m’": "m'", "bh": "buah", "ttk": "titik", "hari.m2": "hari/m2",
                "hari.m4": "hari/m4", "hari.m5": "hari/m5", "hari.m6": "hari/m6", "m 2": "m2"}


def satuan_baku(s: str) -> str:
    s = _teks(s).replace("’", "'").lower()
    return _SATUAN_SAMA.get(s, s)


def baca_induk(wb) -> tuple:
    """Daftar induk sheet "HSPK": ({kode: satuan} item berharga, {bab: judul} untuk bab 1-10)."""
    satuan, bab = {}, {}
    for row in wb["HSPK"].iter_rows(values_only=True):
        row = list(row) + [None] * 10
        no = _teks(row[1]).rstrip(".")
        if not re.match(r"^\d+(\.\d+)*[a-zA-Z]?$", no):
            continue
        if isinstance(row[5], (int, float)) and row[5] > 0:
            satuan[no] = satuan_baku(row[4])
        elif "." not in no and not isinstance(row[1], float):
            bab[no] = _teks(row[3]).title()
    return satuan, {k: v for k, v in bab.items() if int(k) <= BAB_MAKS}


def satuan_dari_uraian(uraian: str) -> str:
    m = _SATUAN_URAIAN.search(uraian)
    return satuan_baku(m.group(1)) if m else ""


def saring_katalog(ahsp: list, satuan_induk: dict) -> list:
    """Analisa sheet utama selalu disimpan (dipakai aplikasi). Analisa sheet lain hanya bila tercantum di
    daftar induk bab 1-10 dan kodenya belum ada (tanpa duplikat "kode@sheet")."""
    hasil, ada = [], set()
    for i in ahsp:
        if i["sheet"] not in SHEET_UTAMA:
            bab = i["kode"].split(".")[0]
            if i["kode"] in ada or i["kode"] not in satuan_induk or not bab.isdigit() or int(bab) > BAB_MAKS:
                continue
        hasil.append(i)
        ada.add(i["kode"])
    return hasil


def tulis_modul(hsd: list, ahsp: list, sumber: str, satuan_induk: dict | None = None, bab: dict | None = None) -> None:
    baris = [
        '"""',
        f"Data HSPK hasil impor otomatis dari: {sumber}",
        "JANGAN diedit manual. Buat ulang dengan:  python tools/impor_hspk.py <file.xlsx>",
        "",
        "HSD  : (tipe, nama, satuan, harga) harga satuan dasar upah / bahan / alat.",
        "AHSP : kode HSPK -> (sheet, uraian, F, [(tipe, nama, satuan, koefisien, harga)])",
        "       F = harga satuan pekerjaan menurut dokumen (sudah termasuk BUK), untuk verifikasi.",
        "SATUAN: kode HSPK -> satuan pekerjaan (daftar induk sheet HSPK).",
        "BAB   : nomor bab -> judul bab (1-10).",
        '"""',
        "",
        f"SUMBER = {sumber!r}",
        "",
        "HSD = [",
    ]
    baris += [f"    {h!r}," for h in hsd]
    baris += ["]", "", "AHSP = {"]
    sudah, satuan = set(), {}
    for i in ahsp:
        kode = i["kode"] if i["kode"] not in sudah else f"{i['kode']}@{i['sheet']}"
        sudah.add(kode)
        satuan[kode] = (satuan_induk or {}).get(i["kode"]) or satuan_dari_uraian(i["uraian"])
        baris.append(f"    {kode!r}: ({i['sheet']!r}, {i['uraian']!r}, {round(i['F'], 4)!r}, [")
        baris += [f"        {k!r}," for k in i["komponen"]]
        baris.append("    ]),")
    baris += ["}", "", "SATUAN = {"]
    baris += [f"    {k!r}: {v!r}," for k, v in satuan.items()]
    baris += ["}", "", "BAB = {"]
    baris += [f"    {k!r}: {v!r}," for k, v in sorted((bab or {}).items(), key=lambda x: int(x[0]))]
    baris += ["}", ""]
    TUJUAN.write_text("\n".join(baris), encoding="utf-8")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", help="workbook HSPK (.xlsx)")
    ap.add_argument("--sumber", default="HSPK Kota Bandung 2027 (update setelah rapat 27 Agustus)")
    a = ap.parse_args()

    print("Membaca workbook ...")
    wb = openpyxl.load_workbook(a.xlsx, read_only=True, data_only=True)
    hsd = baca_hsd(wb)
    satuan_induk, bab = baca_induk(wb)
    ahsp = saring_katalog(baca_ahsp(wb, daftar_sheet(wb)), satuan_induk)

    # verifikasi: koefisien x harga x (1 + BUK) harus sama dengan F di dokumen
    beda = []
    for i in ahsp:
        dasar = sum(k[3] * k[4] for k in i["komponen"])
        buk = i["BUK"] if i["BUK"] is not None else 0.1
        if abs(dasar * (1 + buk) - i["F"]) > max(1.0, 0.002 * i["F"]):
            beda.append(f"{i['sheet']} {i['kode']} {i['uraian'][:60]}")
    tulis_modul(hsd, ahsp, a.sumber, satuan_induk, bab)
    print(f"{len(hsd)} harga dasar, {len(ahsp)} analisa pekerjaan -> {TUJUAN}")
    if beda:
        print(f"{len(beda)} analisa tidak cocok dengan harga F dokumen (periksa sumbernya):")
        for b in beda:
            print("  -", b)
    return 0


if __name__ == "__main__":
    sys.exit(main())
